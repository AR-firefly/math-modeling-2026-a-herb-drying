import json
import os
import sys
import openpyxl as xl
import numpy as np

WORK = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RESULTS = os.path.join(WORK, "results")

_DATA_DIR = os.environ.get("MM_DATA_DIR", WORK)
f1 = os.path.join(_DATA_DIR, "附件", "附件1.xlsx")
f3 = os.path.join(_DATA_DIR, "附件", "附件3") + os.sep


def load_attachment1():
    wb = xl.load_workbook(f1, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    t = np.array([r[0] for r in rows], dtype=float)
    T = np.array([r[1] for r in rows], dtype=float)
    C = np.array([r[2] for r in rows], dtype=float)
    return t, T, C


def repr_list(s):
    return [repr(float(x)) for x in s]


def _attachment1_summary(t, T, C, n, dt, gaps_bad, d, T_max, C_max,
                         i1800, first_dec_t, end_T, end_C, last_10T,
                         last_10C):
    return {
        "文件": "附件1.xlsx",
        "sheet": "Sheet1",
        "表头": ["时间", "温度", "水分浓度"],
        "数据行数": n,
        "时间单位": "s",
        "温度单位": "°C",
        "水分浓度单位": "kg/kg",
        "时间（首值）": repr(float(t[0])),
        "时间（末值）": repr(float(t[-1])),
        "时间步长（唯一值）": [repr(float(x)) for x in np.unique(dt)],
        "是否等间隔60s": bool(np.all(dt == 60.0)),
        "非60s间隔的位置（空 = 无）": [int(i) for i in gaps_bad],
        "重复时间点（空 = 无）": [repr(float(x)) for x in d],
        "温度（全区间最小值）": repr(float(T.min())),
        "温度（全区间最大值）": repr(float(T.max())),
        "温度（最大值对应时间）": repr(float(t[T_max])),
        "水分浓度（全区间最小值）": repr(float(C.min())),
        "水分浓度（全区间最大值）": repr(float(C.max())),
        "水分浓度（最大值对应时间）": repr(float(t[C_max])),
        "温度（首次下降对应时间）": repr(first_dec_t),
        "t1800（温度）": repr(float(T[i1800])),
        "t1800（水分浓度）": repr(float(C[i1800])),
        "末段t>=13000s（温度最小值）": repr(float(end_T.min())),
        "末段t>=13000s（温度最大值）": repr(float(end_T.max())),
        "末段t>=13000s（水分浓度最小值）": repr(float(end_C.min())),
        "末段t>=13000s（水分浓度最大值）": repr(float(end_C.max())),
        "最后10点（温度最小值）": repr(float(last_10T.min())),
        "最后10点（温度最大值）": repr(float(last_10T.max())),
        "最后10点（水分浓度最小值）": repr(float(last_10C.min())),
        "最后10点（水分浓度最大值）": repr(float(last_10C.max())),
        "末段最后10点（温度）": repr_list(last_10T),
        "末段最后10点（水分浓度）": repr_list(last_10C),
        "全量（时间）": repr_list(t),
        "全量（温度）": repr_list(T),
        "全量（水分浓度）": repr_list(C),
    }


def check_attachment1():
    t, T, C = load_attachment1()
    n = len(t)
    dt = np.diff(t)

    gaps_bad = np.where(dt != 60.0)[0]
    u, c = np.unique(t, return_counts=True)
    d = u[c > 1]

    mask_end = t >= 13000.0
    end_T = T[mask_end]
    end_C = C[mask_end]
    last_10T = end_T[-10:]
    last_10C = end_C[-10:]

    dT = np.diff(T)
    neg = np.where(dT < 0.0)[0]
    first_dec_t = t[neg + 1] if len(neg) else None

    T_max = int(np.argmax(T))
    C_max = int(np.argmax(C))

    i1800 = int(np.where(t == 1800.0)[0][0])

    return _attachment1_summary(t, T, C, n, dt, gaps_bad, d, T_max, C_max,
                                i1800, first_dec_t, end_T, end_C, last_10T,
                                last_10C)
    return summary


def check_templates():
    out = {}
    for name in ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"]:
        wb = xl.load_workbook(f3 + name, data_only=True)
        info = {"sheets": list(wb.sheetnames)}
        sheets = {}
        for sn in wb.sheetnames:
            ws = wb[sn]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            sheets[sn] = {
                "维度": ws.dimensions,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "合并单元格数": len(ws.merged_cells.ranges),
                "单元格逐位": rows,
            }
        info["sheets_detail"] = sheets
        out[name] = info
    return out


def main():
    os.makedirs(RESULTS, exist_ok=True)
    att1 = check_attachment1()
    tmpl = check_templates()

    result = {
        "附件1_校验": att1,
        "附件3_模板结构": tmpl,
    }

    out_path = os.path.join(RESULTS, "data_check.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print("附件1：校验结论")
    print("数据行数:", att1["数据行数"], "| 时间范围:",
          att1["时间（首值）"], "~", att1["时间（末值）"], "s")
    print("时间步长唯一值:", att1["时间步长（唯一值）"],
          "| 等间隔60s:", att1["是否等间隔60s"],
          "| 非60s间隔位置:", att1["非60s间隔的位置（空 = 无）"],
          "| 重复时间点:", att1["重复时间点（空 = 无）"])
    print("温度全区间:", att1["温度（全区间最小值）"], "~",
          att1["温度（全区间最大值）"], "°C",
          "| 最大值@t=", att1["温度（最大值对应时间）"], "s")
    print("水分浓度全区间:", att1["水分浓度（全区间最小值）"], "~",
          att1["水分浓度（全区间最大值）"], "kg/kg",
          "| 最大值@t=", att1["水分浓度（最大值对应时间）"], "s")
    print("温度首次下降@t=", att1["温度（首次下降对应时间）"], "s")
    print("t=1800s:", "T=", att1["t1800（温度）"], "°C, C=",
          att1["t1800（水分浓度）"], "kg/kg")
    print("末段(t>=13000s) 温度摆动:",
          att1["末段t>=13000s（温度最小值）"], "~",
          att1["末段t>=13000s（温度最大值）"], "°C")
    print("末段(t>=13000s) 水分浓度摆动:",
          att1["末段t>=13000s（水分浓度最小值）"], "~",
          att1["末段t>=13000s（水分浓度最大值）"], "kg/kg")
    print("最后10点 温度摆动:", att1["最后10点（温度最小值）"], "~",
          att1["最后10点（温度最大值）"], "°C")
    print("最后10点 水分浓度摆动:", att1["最后10点（水分浓度最小值）"], "~",
          att1["最后10点（水分浓度最大值）"], "kg/kg")
    print("[附件3 模板结构]")
    for name, info in tmpl.items():
        print(f"{name}: sheets={info['sheets']}")
        for sn, d in info["sheets_detail"].items():
            print(f"sheet[{sn}] {d['维度']} "
                  f"row={d['max_row']} col={d['max_col']} "
                  f"合并单元格={d['合并单元格数']}")
    print("写入至:", out_path)


if __name__ == "__main__":
    sys.exit(main())
