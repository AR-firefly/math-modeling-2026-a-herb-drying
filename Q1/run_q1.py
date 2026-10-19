from __future__ import annotations
import hashlib
import json
import os
import platform
import sys
import time
from datetime import datetime
from pathlib import Path
import numpy as np

_out = sys.stdout
if hasattr(_out, "reconfigure"):
    _out.reconfigure(encoding="utf-8", errors="replace")

_CODE = Path(os.path.realpath(__file__))
ROOT = _CODE.parents[1]
sys.path.insert(0, str(ROOT / "src"))

from solver_q1 import (
    solve_q1, ATT1, ATT2, PDF, R_CM, RHO_D, HM, T_air_at, C_air_at,
    rho_v_sat, rho_v_inf, default_params,
)
from independent_check_q1_temp import (
    load_oven_temperature, duhamel_temperature,
)

SEED = 42
DR_CM = 0.01
DT_IN = 1.0 / 64
DT_OUT = 1.0
T_END = 1800.0
R_OUT_CM = np.arange(0.0, 2.0 + 1e-9, 0.1)
T_TAB = np.array([100, 300, 600, 900, 1200, 1500, 1800], float)
R_TAB_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])

TEMPERATURE_SOURCE = ("analytic —— Bessel 级数 + Duhamel 卷积 n_roots=320")
MOISTURE_SURFACE_COLUMN = "Richardson p=2.0, N400/N800 (路线①)"

FIELD_T_STRIDE = 20
FIELD_R_STRIDE = 2

OUTPUT_FILES = ["result1.xlsx", "result1_t0_row.xlsx", "result1_caseB.xlsx", "results.json",
                "q1_table1_temperature.csv", "q1_table2_moisture_caseA.csv",
                "q1_surface_column_detail.csv",
                "q1_temperature_numeric_crosscheck.csv", "q1_temperature_analytic_series.csv",
                "q1_field_T_analytic.csv", "q1_field_C_A_numeric.csv",
                "q1_field_C_B_numeric.csv",
                "fig_data/radial_profiles.csv", "fig_data/case_AB_curves.csv"]


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def analytic_T(t_query, r_cm):
    t_air, T_oven = load_oven_temperature()
    return duhamel_temperature(np.asarray(t_query, float),
                               np.asarray(r_cm, float) / 100.0, t_air, T_oven, n_roots=320)


def input_fingerprints():
    inputs = {}
    for p in (ATT1, ATT2, PDF):
        pp = Path(p)
        inputs[pp.name] = {"path": str(pp), "sha256": sha256(pp) if pp.exists() else None,
                           "exists": pp.exists()}
    scripts = {}
    script_paths = {
        "solver_q1.py": ROOT / "Q1" / "solver_q1.py",
        "run_q1.py": ROOT / "Q1" / "run_q1.py",
        "convergence_q1.py": ROOT / "src" / "convergence_q1.py",
        "verification_q1.py": ROOT / "src" / "verification_q1.py",
    }
    for s, sp in script_paths.items():
        if sp.exists():
            scripts[s] = sha256(sp)
    tool_ref = ROOT / "src" / "independent_check_q1_temp.py"
    tool_ref_sha = sha256(tool_ref) if tool_ref.exists() else None
    return inputs, scripts, tool_ref_sha


def run_cases():
    runs = {}
    for case in ("A", "B"):
        o = solve_q1(case=case, dr_cm=DR_CM, dt=DT_IN, t_end=T_END, dt_out=DT_OUT)
        runs[case] = o
        print(f"[run] 口径 {case}: Δr={DR_CM} cm (N={o['N']}), Δt={DT_IN} s, "
              f"步数 {o['n_steps']}, 耗时 {o['elapsed_s']:.1f} s, Picard {o['picard']}")
    return runs["A"], runs["B"]


def surface_sweep(dt):
    surf = {}
    for dr_cm, tag in ((0.01, "N200"), (0.005, "N400"), (0.0025, "N800")):
        o = solve_q1(case="A", dr_cm=dr_cm, dt=dt, t_end=T_END, dt_out=DT_OUT)
        surf[tag] = o["C_out"][1:, -1].copy()
        print(f"[surface] {tag}: 耗时 {o['elapsed_s']:.1f} s  末值 {surf[tag][-1]:.9f}")
    p_surf = 2.0
    surf_ex = surf["N800"] + (surf["N800"] - surf["N400"]) / (2.0 ** p_surf - 1.0)
    surf_ex2 = surf["N400"] + (surf["N400"] - surf["N200"]) / (2.0 ** p_surf - 1.0)
    unc_t = np.maximum(np.abs(surf["N800"] - surf_ex), np.abs(surf_ex2 - surf_ex))
    surf_unc = {tag: float(np.max(np.abs(surf[tag] - surf_ex))) for tag in surf}
    surf_unc["per_time_max"] = float(unc_t.max())
    surf_unc["per_time_at_1800s"] = float(unc_t[-1])
    surf_unc["per_time_max_at_t_s"] = int(np.argmax(unc_t) + 1)
    print(f"[surface] 逐时刻网格不确定度：max {unc_t.max():.3e} @t={int(np.argmax(unc_t))+1} s；"
          f"t=1800 s 时 {unc_t[-1]:.3e} kg/kg")
    return surf, surf_ex, surf_ex2, unc_t, surf_unc


def write_xlsx(path, T_rows, C_rows, times, sheet_names=("温度", "水分浓度")):
    import openpyxl
    wb = openpyxl.Workbook()
    for i, (sn, data) in enumerate(zip(sheet_names, (T_rows, C_rows))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = sn
        ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
        for j, r in enumerate(R_OUT_CM):
            ws.cell(row=1, column=2 + j, value=float(round(r, 1)))
        for k, t in enumerate(times):
            ws.cell(row=2 + k, column=1, value=int(round(t)))
            for j in range(R_OUT_CM.size):
                ws.cell(row=2 + k, column=2 + j, value=float(round(data[k, j], 4)))
    wb.save(path)


def result_tables(A, B, surf, surf_ex, surf_ex2, unc_t, i_out, t_out, n_sel):
    T_win_num = A["T_out"][n_sel][:, i_out]
    T_all_num = np.vstack([A["T_out"][0][i_out][None, :], T_win_num])
    T_all = analytic_T(t_out, R_OUT_CM)
    if T_all.shape != T_all_num.shape:
        raise AssertionError("解析/数值温度场形状不一致")
    T_tab = T_all[n_sel]
    T0_row = T_all[:1]
    T_err_all = np.abs(T_all - T_all_num)

    C_tab = A["C_out"][n_sel][:, i_out].copy()
    C_tab_raw_surface = C_tab[:, -1].copy()
    C_tab[:, -1] = surf_ex
    C_tab_B = B["C_out"][n_sel][:, i_out]
    C0_row = A["C_out"][0][i_out]

    np.savetxt(res_of("q1_surface_column_detail.csv"),
               np.column_stack([t_out[n_sel], surf["N200"], surf["N400"], surf["N800"],
                                surf_ex, surf_ex2, unc_t]),
               delimiter=",", comments="",
               header="t_s,C_surf_N200,C_surf_N400,C_surf_N800,C_surf_Richardson_p2,"
                      "C_surf_Richardson_from_N200N400,grid_uncertainty_conservative")

    write_xlsx(res_of("result1.xlsx"), T_tab, C_tab, t_out[n_sel])
    write_xlsx(res_of("result1_t0_row.xlsx"), T0_row, C0_row[None, :], np.array([0]))
    write_xlsx(res_of("result1_caseB.xlsx"), T_tab, C_tab_B, t_out[n_sel])
    return dict(T_all=T_all, T_all_num=T_all_num, T_tab=T_tab, T0_row=T0_row,
                T_err_all=T_err_all, C_tab=C_tab, C_tab_B=C_tab_B,
                C_tab_raw_surface=C_tab_raw_surface, C0_row=C0_row)


def res_of(name):
    return ROOT / "results" / name


def field_series(A, B, C0_row, surf_ex, t_out, T_all_num):
    i_field_t = np.arange(0, t_out.size, FIELD_T_STRIDE)
    i_field_r = np.arange(0, A["N"] + 1, FIELD_R_STRIDE)
    t_field_s = t_out[i_field_t]
    r_field_cm = A["r"][i_field_r] * 100.0
    if t_field_s.size != 91 or r_field_cm.size != 101:
        raise AssertionError(f"场导出网格不是 91×101（实得 {t_field_s.size}×{r_field_cm.size}）")

    T_field = analytic_T(t_field_s, r_field_cm)
    C_A_field = A["C_out"][np.ix_(i_field_t, i_field_r)].copy()
    C_B_field = B["C_out"][np.ix_(i_field_t, i_field_r)].copy()
    surf_ex_1801 = np.concatenate(([C0_row[-1]], surf_ex))
    if surf_ex_1801.size != t_out.size:
        raise AssertionError("表面列 Richardson 序列长度与输出时刻不符")
    C_A_field[:, -1] = surf_ex_1801[i_field_t]
    T_series = analytic_T(t_out, np.array([0.0, R_CM]))

    hdr_field = "t_s," + ",".join(f"{v:.2f}" for v in r_field_cm)
    np.savetxt(res_of("q1_field_T_analytic.csv"),
               np.column_stack([t_field_s, T_field]), delimiter=",", comments="",
               header=hdr_field)
    np.savetxt(res_of("q1_field_C_A_numeric.csv"),
               np.column_stack([t_field_s, C_A_field]), delimiter=",", comments="",
               header=hdr_field)
    np.savetxt(res_of("q1_field_C_B_numeric.csv"),
               np.column_stack([t_field_s, C_B_field]), delimiter=",", comments="",
               header=hdr_field)
    np.savetxt(res_of("q1_temperature_analytic_series.csv"),
               np.column_stack([t_out, T_series]), delimiter=",", comments="",
               header="t_s,T_center_C,T_surface_C")
    np.savetxt(res_of("q1_temperature_numeric_crosscheck.csv"),
               np.column_stack([t_out, T_all_num]), delimiter=",", comments="",
               header="t_s," + ",".join(f"{rc:g}" for rc in R_OUT_CM))
    return i_field_t, i_field_r, t_field_s, r_field_cm, T_field


def error_crosscheck(A, t_out, T_err_all, T_field, i_field_t, i_field_r,
                     t_field_s, r_field_cm):
    i_win = t_out >= 1.0
    err_win = T_err_all[i_win]
    k_win = np.unravel_index(int(np.argmax(err_win)), err_win.shape)
    err_field = np.abs(T_field - A["T_out"][np.ix_(i_field_t, i_field_r)])
    k_fld = np.unravel_index(int(np.argmax(err_field)), err_field.shape)
    return {
        "note": "数值解温度相对解析解的逐点偏差",
        "numeric_file": "results/q1_temperature_numeric_crosscheck.csv",
        "window": "t=1…1800 s × r=0…2.0 cm",
        "max_abs_err_window_C": float(err_win.max()),
        "max_abs_err_window_at": {"t_s": float(t_out[i_win][k_win[0]]),
                                  "r_cm": float(R_OUT_CM[k_win[1]])},
        "rms_abs_err_window_C": float(np.sqrt(np.mean(err_win ** 2))),
        "max_abs_err_at_1800s_C": float(err_win[-1].max()),
        "max_abs_err_incl_t0_C": float(T_err_all.max()),
        "max_abs_err_field_91x101_C": float(err_field.max()),
        "max_abs_err_field_at": {"t_s": float(t_field_s[k_fld[0]]),
                                 "r_cm": float(r_field_cm[k_fld[1]])},
        "acceptance_C": 5.0e-05,
        "meets_acceptance_full_window": bool(err_win.max() < 5.0e-05),
    }, err_win, k_win


def analytic_benchmarks(A, i_out):
    T_an_1800 = analytic_T([T_END], R_OUT_CM)[0]
    T_an_tab = analytic_T(T_TAB, R_TAB_CM)
    i_tab = [int(round(rc / DR_CM)) for rc in R_TAB_CM]
    j_tab = [int(round(rc / 0.1)) for rc in R_TAB_CM]
    T_num_tab = np.array([[A["T_out"][int(t)][i] for i in i_tab] for t in T_TAB])
    err_tab = np.abs(T_num_tab - T_an_tab)
    err_1800 = np.abs(A["T_out"][int(T_END)][i_out] - T_an_1800)
    return T_an_1800, T_an_tab, i_tab, j_tab, T_num_tab, err_tab, err_1800


def conservation(A, B):
    sc = A["scalars"]
    cons_rel = float(abs((sc["mass"][0] - sc["mass"][-1]) - sc["cum_flux"][-1])
                     / (sc["mass"][0] - sc["mass"][-1]))
    scB = B["scalars"]
    cons_rel_B = float(abs((scB["mass"][0] - scB["mass"][-1]) - scB["cum_flux"][-1])
                       / max(abs(scB["mass"][0] - scB["mass"][-1]), 1e-300))
    return sc, scB, cons_rel, cons_rel_B


def _summary_meta(inputs, tool_ref_sha):
    return {
        "meta": {
            "question": "Q1 预热平衡阶段（0–1800 s）",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "seed": SEED,
            "config": {"dr_cm": DR_CM, "dt_internal_s": DT_IN, "dt_out_s": DT_OUT,
                       "t_end_s": T_END, "picard_tol": 1e-12, "picard_max": 50,
                       "interface_D_rule": "arithmetic mean of D(C_i), D(C_i+1)",
                       "surface_column": "Richardson p=2.0, N400/N800 (路线①)",
                       "temperature_source": TEMPERATURE_SOURCE,
                       "moisture_surface_column": MOISTURE_SURFACE_COLUMN},
            "python": platform.python_version(),
            "numpy": np.__version__,
        },
        "inputs_sha256": inputs,
        "reference_tool_sha256": tool_ref_sha,
        "units": {"T": "℃", "C": "kg/kg(干基)", "r": "cm", "t": "s",
                  "mass": "kg water per kg dry solids, 归一化质量比"},
    }


def _summary_temperature(T_all, T_tab, T_an_1800, T_an_tab, err_tab,
                         err_1800, T_num_tab, crosscheck):
    return {
        "temperature": {
            "source": TEMPERATURE_SOURCE,
            "delivered_is": "analytic（解析解）——见analytic块",
            "reported": T_tab[-1].round(6).tolist(),
            "analytic": {
                "note": ("常数物性下该线性热传导问题的精确解，"
                         "Bessel 级数 + Duhamel 卷积 n_roots=320；实现对 "
                         "src/independent_check_q1_temp.py。"
                         "result1*.xlsx 的「温度」sheet 四位小数取自本块。"),
                "t_s": float(T_END),
                "T_at_1800s_r_cm": R_OUT_CM.tolist(),
                "T_at_1800s": T_all[-1].round(9).tolist(),
                "table1": T_an_tab.round(9).tolist(),
                "anchors_1800s_C": {"center": float(T_all[-1, 0]),
                                    "surface": float(T_all[-1, -1])},
            },
            "numeric_crosscheck": crosscheck,
            "analytic_1800s": T_an_1800.round(9).tolist(),
            "abs_err_1800s": err_1800.round(12).tolist(),
            "max_abs_err_1800s": float(err_1800.max()),
            "table1_numeric": T_num_tab.round(6).tolist(),
            "table1_analytic": T_an_tab.round(9).tolist(),
            "table1_max_abs_err": float(err_tab.max()),
            "error_vs_analytic_source": "src/independent_check_q1_temp.py n_roots=320",
        },
    }


def _summary_moisture(A, B, surf, surf_ex, surf_unc, C_tab,
                       C_tab_raw_surface, C_tab_B, sc, scB, cons_rel,
                       cons_rel_B, i_tab, j_tab, idx):
    return {
        "moisture_caseA": {
            "surface_C_1800s_gridconverged": float(surf_ex[-1]),
            "surface_C_1800s_raw_N200": float(C_tab_raw_surface[-1]),
            "surface_C_1800s_raw_N400": float(surf["N400"][-1]),
            "surface_C_1800s_raw_N800": float(surf["N800"][-1]),
            "surface_grid_uncertainty": surf_unc,
            "center_C_1800s": float(A["C_out"][int(T_END)][0]),
            "mass_ratio_1800s": float(sc["mass"][-1] / sc["mass"][0]),
            "mass_ratio_lost_pct": float(100.0 * (1.0 - sc["mass"][-1] / sc["mass"][0])),
            "cumulative_flux": float(sc["cum_flux"][-1]),
            "conservation_rel_error": cons_rel,
            "center_value_dev_from_init": float(abs(A["C_out"][int(T_END)][0] - 2.55)),
        },
        "moisture_caseB": {
            "note": "含题目外假设 B1与 B2；",
            "rho_d_kg_m3": RHO_D,
            "surface_C_1800s": float(B["C_out"][int(T_END)][-1]),
            "center_C_1800s": float(B["C_out"][int(T_END)][0]),
            "mass_ratio_1800s": float(scB["mass"][-1] / scB["mass"][0]),
            "conservation_rel_error": cons_rel_B,
            "scoping_reference": {"surface_C_quasi_steady": 2.549885,
                                  "flux_kg_m2_s": 6.534236e-09},
        },
        "table2_caseA": {"times_s": T_TAB.tolist(), "radii_cm": R_TAB_CM.tolist(),
                         "note": "与 result1.xlsx 同源",
                         "C": np.array([[C_tab[idx(t) - 1][j] for j in j_tab]
                                        for t in T_TAB]).round(6).tolist()},
        "table2_caseB": {"times_s": T_TAB.tolist(), "radii_cm": R_TAB_CM.tolist(),
                         "C": np.array([[B["C_out"][idx(t)][i] for i in i_tab]
                                        for t in T_TAB]).round(9).tolist()},
        "picard": {"caseA": A["picard"], "caseB": B["picard"]},
    }


def build_summary(A, B, t_out, surf, surf_ex, surf_unc, T_all, T_tab,
                  T_an_1800, T_an_tab, err_tab, err_1800, T_num_tab,
                  C_tab, C_tab_raw_surface, C_tab_B, crosscheck, sc, scB,
                  cons_rel, cons_rel_B, i_out, i_tab, j_tab, idx,
                  inputs, tool_ref_sha):
    summary = {}
    summary.update(_summary_meta(inputs, tool_ref_sha))
    summary.update(_summary_temperature(T_all, T_tab, T_an_1800, T_an_tab,
                                        err_tab, err_1800, T_num_tab,
                                        crosscheck))
    summary.update(_summary_moisture(A, B, surf, surf_ex, surf_unc, C_tab,
                                     C_tab_raw_surface, C_tab_B, sc, scB,
                                     cons_rel, cons_rel_B, i_tab, j_tab,
                                     idx))
    return summary


def write_table_csv(T_an_tab, C_tab, idx, j_tab):
    hdr = "时间/s," + ",".join(f"{r:g}" for r in R_TAB_CM)
    lines = [hdr]
    for i, t in enumerate(T_TAB):
        lines.append(f"{t:g}," + ",".join(f"{v:.4f}" for v in T_an_tab[i]))
    res_of("q1_table1_temperature.csv").write_text("\n".join(lines), encoding="utf-8-sig")
    lines = [hdr]
    for i, t in enumerate(T_TAB):
        lines.append(f"{t:g}," + ",".join(
            f"{C_tab[idx(t) - 1][j]:.4f}" for j in j_tab))
    res_of("q1_table2_moisture_caseA.csv").write_text("\n".join(lines), encoding="utf-8-sig")


def write_figure_data(A, B, i_out, idx):
    figd = ROOT / "results" / "fig_data"
    prof_t = T_TAB
    np.savetxt(figd / "radial_profiles.csv",
               np.column_stack([np.repeat(prof_t, len(R_OUT_CM)),
                                np.tile(R_OUT_CM, len(prof_t)),
                                np.array([A["T_out"][idx(t)][i_out] for t in prof_t]).ravel(),
                                np.array([A["C_out"][idx(t)][i_out] for t in prof_t]).ravel(),
                                np.array([B["C_out"][idx(t)][i_out] for t in prof_t]).ravel()]),
               delimiter=",", header="t_s,r_cm,T_C_CaseA,C_CaseA,C_CaseB", comments="")

    ts = A["t_out"][1:]
    w_all = np.array([C_air_at(t) for t in ts])
    Ta_all = np.array([T_air_at(t) for t in ts])
    CsA = A["C_out"][1:, -1]
    CsB = B["C_out"][1:, -1]
    TsB = B["T_out"][1:, -1]
    qA = HM * (CsA - w_all)
    qB = HM * (rho_v_sat(TsB) - rho_v_inf(w_all, Ta_all)) / RHO_D
    np.savetxt(figd / "case_AB_curves.csv",
               np.column_stack([ts, CsA, CsB, A["C_out"][1:, 0], B["C_out"][1:, 0],
                                qA * RHO_D, qB * RHO_D]),
               delimiter=",",
               header="t_s,C_surf_CaseA,C_surf_CaseB,C_center_CaseA,C_center_CaseB,"
                      "flux_CaseA_kg_m2_s,flux_CaseB_kg_m2_s", comments="")

    (figd / "README.json").write_text(json.dumps({
        "note": "作图数据入口",
        "files": {
            "radial_profiles.csv":
                "列：t_s, r_cm, T_C_CaseA, C_CaseA, C_CaseB（t=100/300/600/900/1200/1500/1800 s）",
            "case_AB_curves.csv":
                "列：t_s, C_surf_CaseA, C_surf_CaseB, C_center_CaseA, C_center_CaseB, "
                "flux_CaseA_kg_m2_s, flux_CaseB_kg_m2_s（t=1…1800 s 逐秒）",
            "error_bars.json": "网格/时间收敛偏差",
            "sensitivity.csv": "单/多参数灵敏度扫描",
            "envelope_hm.csv": "h_m 驱动势包络扫描",
            "ab_compare.json": "口径 A/B 对比关键数值",
        },
    }, ensure_ascii=False, indent=2), encoding="utf-8")


def build_uncertainty(res, surf_unc, summary):
    unc = {
        "note": "A10.1 要求的不确定度标注。表面列分时刻值的逐秒完整序列见 "
                "results/q1_surface_column_detail.csv 的 grid_uncertainty_conservative 列。",
        "surface_column_moisture_grid": {
            "note": "本部分含两个不同对象，不可混读；均为网格不确定度",
            "method": "Richardson p=2.00；保守值再与 N=200/N=400 外推取较大",
            "delivered_column_conservative": {
                "definition": "交付表面列"
                              "max(|C_N800−C_ex|, |C_{N200,N400}外推−C_ex|)。",
                "per_time_max": surf_unc["per_time_max"],
                "per_time_max_at_t_s": surf_unc["per_time_max_at_t_s"],
                "per_time_at_1800s": surf_unc["per_time_at_1800s"],
                "per_time_full_series": "results/q1_surface_column_detail.csv  grid_uncertainty_conservative",
            },
            "raw_node_deviation": {
                "definition": "未外推的原始节点值与收敛值之差，",
                "max_over_window": {"N200": surf_unc["N200"], "N400": surf_unc["N400"],
                                    "N800": surf_unc["N800"]},
                "at_1800s": {},
            },
        },
    }

    _cv = res / "q1_convergence.json"
    if _cv.exists():
        _t4 = json.loads(_cv.read_text(encoding="utf-8"))["table4_dr_sweep_moisture_surface"]
        unc["surface_column_moisture_grid"]["raw_node_deviation"]["at_1800s"] = dict(
            _t4["grid_uncertainty_vs_richardson"])
        unc["surface_column_moisture_grid"]["raw_node_deviation"]["at_1800s_note"] = \
            "源 results/q1_convergence.json  table4_dr_sweep_moisture_surface.grid_uncertainty_vs_richardson"
    _we = res / "q1_window_error.json"
    if _we.exists():
        _w = json.loads(_we.read_text(encoding="utf-8"))["moisture_window_error"]
        unc["moisture_time_step_residual"] = {
            "delivered_dt_s": _w["delivered_dt_s"],
            "error_estimate_p1": _w["delivered_dt_error_estimate_p1"],
            "criterion_C_analog": _w["criterion_C_analog"],
            "n_cells_over_criterion_p1": _w["n_cells_over_criterion_p1"],
            "n_cells_total": _w["n_cells_total"],
            "note": "Δt=1/64 s 自身时间离散残差估计（p=1）；超判据的 2 格均位于早期表面列，"
                    "该处网格不确定度（2.71e-4）比时间误差大 3.3 倍为主导项",
            "source": "results/q1_window_error.json（由 src/window_error_q1.py 生成）",
        }
    else:
        unc["moisture_time_step_residual"] = {
            "source": "results/q1_window_error.json",
        }
    unc["temperature_window_error_crosscheck"] = {
        "value_C": summary["temperature"]["numeric_crosscheck"]["max_abs_err_window_C"],
        "at": summary["temperature"]["numeric_crosscheck"]["max_abs_err_window_at"],
        "source": "results.json  temperature.numeric_crosscheck",
    }
    return unc


def _run_manifest(inputs, scripts, tool_ref_sha, summary):
    manifest = {
        "seed": SEED,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "inputs_sha256": inputs,
        "reference_tool_sha256": tool_ref_sha,
        "script_sha256": scripts,
        "config": summary["meta"]["config"],
        "wall_s": summary["wall_s"],
        "outputs": {},
        "uncertainty": {},
    }
    return manifest


def _print_summary(res, t_out, err_win, k_win, err_1800, surf_ex,
                   C_tab_raw_surface, surf_unc, sc, scB, cons_rel, B,
                   summary):
    _win = t_out >= 1.0
    print(f"\n[summary] 温度 = 解析解")
    print(f"[summary] 数值解 vs 解析解 全窗口最大偏差 = {err_win.max():.3e} ℃ "
          f"@t={t_out[_win][k_win[0]]:.0f} s, r={R_OUT_CM[k_win[1]]:g} cm "
          f"（验收线 5e-5 ℃，{'达标' if err_win.max() < 5e-5 else '超标'}）；"
          f"t=1800 s 处 {err_1800.max():.3e} ℃")
    print(f"[summary] 温度 t=1800 s 与解析解最大偏差 = {err_1800.max():.3e} ℃")
    print(f"[summary] 生产场落盘：91×101（t 每 20 s、r 每 0.02 cm） q1_field_T_analytic.csv / "
          f"q1_field_C_A_numeric.csv / q1_field_C_B_numeric.csv")
    print(f"[summary] 口径A 表面 C(1800) 交付值 = {surf_ex[-1]:.6f}（Richardson 外推）")
    print(f"[summary]   └ 未外推的 N=200 原始节点值 {C_tab_raw_surface[-1]:.6f}，"
          f"其与收敛值之差 {surf_unc['N200']:.2e}")
    print(f"[summary] 口径A 总水量比 = {sc['mass'][-1]/sc['mass'][0]:.9f}；守恒相对误差 {cons_rel:.3e}")
    print(f"[summary] 口径B 表面 C(1800) = {B['C_out'][int(T_END)][-1]:.9f}；"
          f"总水量比 {scB['mass'][-1]/scB['mass'][0]:.9f}")
    print(f"[summary] 总耗时 {summary['wall_s']:.1f} s；已写 {res/'result1.xlsx'} 等")


def main():
    t0 = time.time()
    res = ROOT / "results"
    (res / "fig_data").mkdir(parents=True, exist_ok=True)

    inputs, scripts, tool_ref_sha = input_fingerprints()
    A, B = run_cases()
    t_out = A["t_out"]
    n_sel = np.arange(1, t_out.size)
    surf, surf_ex, surf_ex2, unc_t, surf_unc = surface_sweep(DT_IN)

    i_out = np.array([int(round(rc / DR_CM)) for rc in R_OUT_CM])
    tab = result_tables(A, B, surf, surf_ex, surf_ex2, unc_t, i_out, t_out, n_sel)
    T_all, T_all_num = tab["T_all"], tab["T_all_num"]
    T_tab, T_err_all, C_tab = tab["T_tab"], tab["T_err_all"], tab["C_tab"]
    C_tab_B = tab["C_tab_B"]
    C_tab_raw_surface, C0_row = tab["C_tab_raw_surface"], tab["C0_row"]

    i_field_t, i_field_r, t_field_s, r_field_cm, T_field = field_series(
        A, B, C0_row, surf_ex, t_out, T_all_num)
    crosscheck, err_win, k_win = error_crosscheck(
        A, t_out, T_err_all, T_field, i_field_t, i_field_r, t_field_s, r_field_cm)
    T_an_1800, T_an_tab, i_tab, j_tab, T_num_tab, err_tab, err_1800 = \
        analytic_benchmarks(A, i_out)
    sc, scB, cons_rel, cons_rel_B = conservation(A, B)

    def idx(t):
        return int(t)

    summary = build_summary(A, B, t_out, surf, surf_ex, surf_unc, T_all, T_tab,
                            T_an_1800, T_an_tab, err_tab, err_1800, T_num_tab,
                            C_tab, C_tab_raw_surface, C_tab_B, crosscheck, sc, scB,
                            cons_rel, cons_rel_B, i_out, i_tab, j_tab, idx,
                            inputs, tool_ref_sha)
    summary["wall_s"] = time.time() - t0
    (res / "results.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    manifest = _run_manifest(inputs, scripts, tool_ref_sha, summary)

    write_table_csv(T_an_tab, C_tab, idx, j_tab)
    write_figure_data(A, B, i_out, idx)

    for f in OUTPUT_FILES:
        p = res / f
        manifest["outputs"][f] = sha256(p) if p.exists() else None

    manifest["uncertainty"] = build_uncertainty(res, surf_unc, summary)

    (res / "run_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    _print_summary(res, t_out, err_win, k_win, err_1800, surf_ex,
                   C_tab_raw_surface, surf_unc, sc, scB, cons_rel, B,
                   summary)


if __name__ == "__main__":
    sys.exit(main())
