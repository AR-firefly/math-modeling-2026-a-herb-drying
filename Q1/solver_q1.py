from __future__ import annotations
import hashlib
import json
import os
import time
from pathlib import Path
import numpy as np
from scipy.linalg import solve_banded
import openpyxl as xl

RHO = 820.0
CP = 2600.0
K = 0.36
H = 25.0
HM = 8.0e-7
D0 = 7.0e-9
BETA = 0.89

R_CM = 2.0
LENGTH_CM = 25.0
R = R_CM / 100.0

T_INIT = 28.0
C_INIT = 2.55

P_ATM = 101325.0
R_DA = 287.05
R_V = 461.5
EPS_RATIO = 0.622
RHO_D = RHO / (1.0 + C_INIT)

ROOT = Path(os.path.realpath(__file__)).parents[1]

_DATA_DIR = os.environ.get("MM_DATA_DIR", str(ROOT))
ATT1 = str(Path(_DATA_DIR) / "附件" / "附件1.xlsx")
ATT2 = str(Path(_DATA_DIR) / "附件" / "附件2.xlsx")
PDF = str(Path(_DATA_DIR) / "A题.pdf")

_ATT1_CACHE = None


def att1():
    global _ATT1_CACHE
    if _ATT1_CACHE is None:
        wb = xl.load_workbook(ATT1, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        t = np.array([r[0] for r in rows], float)
        T = np.array([r[1] for r in rows], float)
        C = np.array([r[2] for r in rows], float)
        _ATT1_CACHE = (t, T, C)
    return _ATT1_CACHE


def T_air_at(t):
    t0, T0, _ = att1()
    return np.interp(t, t0, T0)


def C_air_at(t):
    t0, _, C0 = att1()
    return np.interp(t, t0, C0)


def D_of_C(C):
    return D0 * np.exp(-BETA / C)


def D_safe(C):
    return D0 * np.exp(-BETA / np.maximum(C, 1e-12))


def alpha():
    return K / (RHO * CP)


def sat_pressure_pa(T_C):
    return 611.21 * np.exp(17.502 * T_C / (240.97 + T_C))


def rho_v_sat(T_C):
    T_K = T_C + 273.15
    return sat_pressure_pa(T_C) / (R_V * T_K)


def rho_v_inf(w, T_C):
    T_K = T_C + 273.15
    p_v = w * P_ATM / (EPS_RATIO + w)
    p_da = P_ATM - p_v
    rho_da = p_da / (R_DA * T_K)
    return w * rho_da


def build_grid(dr_cm):
    dr = dr_cm / 100.0
    N = int(round(R / dr))
    if abs(N * dr - R) > 1e-12:
        raise ValueError(f"dr={dr_cm} cm 不能整除 R={R_CM} cm")
    r = np.arange(N + 1) * dr
    return r, dr, N


def heat_step(T, r, dr, dt, T_air):
    N = len(r) - 1
    al = alpha()
    lam = dt * al / (dr * dr)

    d = np.zeros(N + 1)
    u = np.zeros(N)
    l = np.zeros(N + 1)
    rhs = np.zeros(N + 1)

    rp = r + dr / 2.0
    rm = r - dr / 2.0

    for i in range(1, N):
        d[i] = 1.0 + lam * (rp[i] + rm[i]) / r[i]
        l[i] = -lam * rm[i] / r[i]
        u[i] = -lam * rp[i] / r[i]
        rhs[i] = T[i]

    d[0] = 1.0 + 4.0 * lam
    u[0] = -4.0 * lam
    rhs[0] = T[0]

    m = R * dr - dr * dr / 4.0
    a_in = 2.0 * K * (R - dr / 2.0) / dr
    a_out = 2.0 * H * R
    G = RHO * CP * m / dt
    d[N] = 1.0 + (a_in + a_out) / G
    l[N] = -a_in / G
    rhs[N] = T[N] + (a_out / G) * T_air

    ab = np.zeros((3, N + 1))
    ab[0, 1:] = u
    ab[1, :] = d
    ab[2, :-1] = l[1:]
    return solve_banded((1, 1), ab, rhs)


def solve_heat_constant(dr_cm=0.005, dt=1.0, t_end=1800.0, record_every=1.0):
    r, dr, N = build_grid(dr_cm)
    T = np.full(N + 1, T_INIT)

    n_steps = int(round(t_end / dt))
    rec_gap = max(1, int(round(record_every / dt)))
    t_out = []
    T_hist = []

    t = 0.0
    t_out.append(t)
    T_hist.append(T.copy())
    for step in range(1, n_steps + 1):
        t = step * dt
        T_air = T_air_at(t)
        T = heat_step(T, r, dr, dt, T_air)
        if step % rec_gap == 0:
            t_out.append(t)
            T_hist.append(T.copy())
    return np.array(t_out), np.array(T_hist), r


def surface_robin_terms(dr, gam_face, h_coef):
    m_over_pi = R * dr - dr * dr / 4.0
    a_in = 2.0 * (R - dr / 2.0) * gam_face / dr
    a_out = 2.0 * R * h_coef
    return m_over_pi, a_in, a_out


def interp_driver(t, kind):
    t0, T0, C0 = att1()
    if kind == "linear":
        return float(np.interp(t, t0, T0)), float(np.interp(t, t0, C0))
    if kind == "step":
        i = int(np.searchsorted(t0, t, side="right")) - 1
        i = max(0, min(i, t0.size - 1))
        return float(T0[i]), float(C0[i])
    raise ValueError(f"未知插值方式 kind={kind!r}")


def _mass_step_prep(C_old, C_it, r, dr, dt, P):
    N = len(r) - 1
    if P is None:
        P = default_params()
    C = C_it.copy()
    rp = r + dr / 2.0
    rm = r - dr / 2.0
    rho_d = P["RHO"] / (1.0 + P["C_INIT"])

    geo_in = np.zeros(N + 1)
    geo_out = np.zeros(N + 1)
    with np.errstate(divide="ignore", invalid="ignore"):
        geo_in[1:N] = dt * rm[1:N] / (r[1:N] * dr * dr)
        geo_out[1:N] = dt * rp[1:N] / (r[1:N] * dr * dr)
    c0_factor = 4.0 * dt / (dr * dr)

    rhs_old = C_old.copy()

    converged = False
    n_iter = 0
    d = np.zeros(N + 1)
    l = np.zeros(N + 1)
    u = np.zeros(N)
    ab = np.zeros((3, N + 1))
    return (N, P, C, rho_d, geo_in, geo_out, c0_factor, rhs_old, converged,
            n_iter, d, l, u, ab)


def _picard_solve(C, N, dr, dt, P, picard_tol, picard_max, d, l, u, ab, geo_in, geo_out, c0_factor, rho_d, case, T_surf_new, T_air_new, C_air_new, rhs_old, converged, n_iter):
    for k in range(1, picard_max + 1):
        n_iter = k
        Cs = np.maximum(C, 1e-12)
        if P["D_CONST"] is None:
            Dv = P["D0"] * np.exp(-P["BETA"] / Cs)
        else:
            Dv = np.full(C.shape, P["D_CONST"])
        face = P["D_FACE"]
        if face == "arith":
            dm = 0.5 * (Dv[:-1] + Dv[1:])
        elif face == "left":
            dm = Dv[:-1].copy()
        elif face == "right":
            dm = Dv[1:].copy()
        elif face == "harm":
            dm = 2.0 * Dv[:-1] * Dv[1:] / (Dv[:-1] + Dv[1:])
        elif face == "faceC":
            Cm = 0.5 * (Cs[:-1] + Cs[1:])
            if P["D_CONST"] is None:
                dm = P["D0"] * np.exp(-P["BETA"] / Cm)
            else:
                dm = np.full(N, P["D_CONST"])
        else:
            raise ValueError(f"未知界面系数规则 D_FACE={face!r}")
        d[1:N] = 1.0 + geo_in[1:N] * dm[0:N - 1] + geo_out[1:N] * dm[1:N]
        l[1:N] = -geo_in[1:N] * dm[0:N - 1]
        u[1:N] = -geo_out[1:N] * dm[1:N]
        a0 = c0_factor * dm[0]
        d[0] = 1.0 + a0
        u[0] = -a0
        m_over_pi, a_in, a_out = surface_robin_terms(dr, dm[N - 1], P["HM"])
        d[N] = m_over_pi / dt + a_in
        l[N] = -a_in
        rhs = np.empty(N + 1)
        rhs[0:N] = rhs_old[0:N]
        rhs[N] = (m_over_pi / dt) * rhs_old[N]
        if P["HM_INF"]:
            d[N] = 1.0
            l[N] = 0.0
            rhs[N] = C_air_new
        elif case == "A":
            d[N] += a_out
            rhs[N] += a_out * C_air_new
        elif case == "B":
            q_B = P["HM"] * (rho_v_sat(T_surf_new) - rho_v_inf(C_air_new, T_air_new)) / rho_d
            rhs[N] -= 2.0 * R * q_B
        else:
            raise ValueError(f"未知口径 case={case!r}（应为 'A' 或 'B'）")
        ab[0, 1:] = u
        ab[1, :] = d
        ab[2, :-1] = l[1:]
        C_new = solve_banded((1, 1), ab, rhs)
        delta = float(np.max(np.abs(C_new - C)))
        C = C_new
        if delta < picard_tol:
            converged = True
            break
    return C, n_iter, converged


def mass_step(C_old, C_it, r, dr, dt, t_new, case="A", T_surf_new=None,
               picard_tol=1e-12, picard_max=50, P=None,
               T_air_new=None, C_air_new=None):
    (N, P, C, rho_d, geo_in, geo_out, c0_factor, rhs_old, converged,
     n_iter, d, l, u, ab) = _mass_step_prep(C_old, C_it, r, dr, dt, P)

    return _picard_solve(C, N, dr, dt, P, picard_tol, picard_max, d, l, u, ab,
                         geo_in, geo_out, c0_factor, rho_d, case, T_surf_new,
                         T_air_new, C_air_new, rhs_old, converged, n_iter)


def heat_matrix(r, dr, dt, h=H, k=K, rho=RHO, cp=CP):
    N = len(r) - 1
    al = k / (rho * cp)
    lam = dt * al / (dr * dr)
    rp = r + dr / 2.0
    rm = r - dr / 2.0
    d = np.zeros(N + 1)
    l = np.zeros(N + 1)
    u = np.zeros(N)
    d[1:N] = 1.0 + lam * (rp[1:N] + rm[1:N]) / r[1:N]
    l[1:N] = -lam * rm[1:N] / r[1:N]
    u[1:N] = -lam * rp[1:N] / r[1:N]
    d[0] = 1.0 + 4.0 * lam
    u[0] = -4.0 * lam
    m_over_pi, a_in, a_out = surface_robin_terms(dr, al, h / (rho * cp))
    G = m_over_pi / dt
    d[N] = 1.0 + (a_in + a_out) / G
    l[N] = -a_in / G
    a_out_over_G = a_out / G
    ab = np.zeros((3, N + 1))
    ab[0, 1:] = u
    ab[1, :] = d
    ab[2, :-1] = l[1:]
    return ab, a_out_over_G


def default_params():
    return dict(H=H, HM=HM, K=K, RHO=RHO, CP=CP, D0=D0, BETA=BETA,
                C_INIT=C_INIT, T_INIT=T_INIT, INTERP="linear",
                D_CONST=None, HM_INF=False, D_FACE="arith")


def _q1_setup(P, dr_cm, dt, t_end, dt_out, record_fields):
    r, dr, N = build_grid(dr_cm)
    n_steps = int(round(t_end / dt))
    if abs(n_steps * dt - t_end) > 1e-9:
        raise ValueError(f"t_end={t_end} 不是 dt={dt} 的整数倍")
    rec_gap = int(round(dt_out / dt))
    if rec_gap < 1 or abs(rec_gap * dt - dt_out) > 1e-9:
        raise ValueError(f"dt_out={dt_out} 必须是 dt={dt} 的整数倍")

    T0 = np.empty(n_steps + 1)
    TN = np.empty(n_steps + 1)
    C0 = np.empty(n_steps + 1)
    CN = np.empty(n_steps + 1)
    MASS = np.empty(n_steps + 1)
    FLUX = np.empty(n_steps + 1)
    TVOL = np.zeros(N + 1)
    TVOL[0] = dr * dr / 4.0
    TVOL[1:N] = 2.0 * r[1:N] * dr
    TVOL[N] = R * dr - dr * dr / 4.0

    T = np.full(N + 1, P["T_INIT"])
    C = np.full(N + 1, P["C_INIT"])
    T0[0], TN[0], C0[0], CN[0] = T[0], T[N], C[0], C[N]
    MASS[0] = float(np.dot(TVOL, C))
    FLUX[0] = 0.0

    ab_heat, a_out_over_G = heat_matrix(r, dr, dt, P["H"], P["K"], P["RHO"], P["CP"])

    if record_fields:
        T_out = np.empty((n_steps // rec_gap + 1, N + 1))
        C_out = np.empty((n_steps // rec_gap + 1, N + 1))
        T_out[0] = T
        C_out[0] = C
    else:
        T_out = C_out = None

    picard_iters = np.zeros(n_steps + 1, dtype=int)
    n_not_conv = 0
    rhs_h = np.empty(N + 1)
    return (r, dr, N, n_steps, rec_gap, T0, TN, C0, CN, MASS, FLUX, TVOL, T,
            C, T_out, C_out, ab_heat, a_out_over_G, picard_iters, n_not_conv,
            rhs_h)


def solve_q1(case="A", dr_cm=0.01, dt=1.0, t_end=1800.0, dt_out=1.0,
             picard_tol=1e-12, picard_max=50, record_fields=True,
             return_scalars=True, **overrides):
    t_start = time.time()
    P = default_params()
    P.update(overrides)
    (r, dr, N, n_steps, rec_gap, T0, TN, C0, CN, MASS, FLUX, TVOL, T, C,
     T_out, C_out, ab_heat, a_out_over_G, picard_iters, n_not_conv,
     rhs_h) = _q1_setup(P, dr_cm, dt, t_end, dt_out, record_fields)

    for step in range(1, n_steps + 1):
        t_new = step * dt
        T_air_new, C_air_new = interp_driver(t_new, P["INTERP"])

        rhs_h[:] = T
        rhs_h[N] += a_out_over_G * T_air_new
        T = solve_banded((1, 1), ab_heat, rhs_h)

        C, k_it, ok = mass_step(C, C, r, dr, dt, t_new, case=case,
                                 T_surf_new=T[N], picard_tol=picard_tol,
                                 picard_max=picard_max, P=P,
                                 T_air_new=T_air_new, C_air_new=C_air_new)
        picard_iters[step] = k_it
        if not ok:
            n_not_conv += 1

        T0[step], TN[step], C0[step], CN[step] = T[0], T[N], C[0], C[N]
        MASS[step] = float(np.dot(TVOL, C))
        if P["HM_INF"]:
            q = 0.0
        elif case == "A":
            q = P["HM"] * (CN[step] - C_air_new)
        else:
            q = (P["HM"] * (rho_v_sat(T[N]) - rho_v_inf(C_air_new, T_air_new))
                 / (P["RHO"] / (1.0 + P["C_INIT"])))
        FLUX[step] = FLUX[step - 1] + 2.0 * R * q * dt

        if record_fields and step % rec_gap == 0:
            i = step // rec_gap
            T_out[i] = T
            C_out[i] = C

    t_out = np.arange(0, n_steps + 1, rec_gap) * dt
    out = dict(
        params=P,
        r=r, dr=dr, N=N, case=case, dt=dt, dt_out=dt_out, t_end=t_end,
        t_out=t_out, T_out=T_out, C_out=C_out,
        scalars=dict(t=np.arange(n_steps + 1) * dt, T0=T0, TN=TN, C0=C0, CN=CN,
                     mass=MASS, cum_flux=FLUX, tvol=TVOL),
        picard=dict(max_iters=int(picard_iters.max()), n_not_converged=int(n_not_conv),
                    mean_iters=float(picard_iters[1:].mean()) if n_steps else 0.0),
        elapsed_s=time.time() - t_start,
        n_steps=n_steps,
    )
    return out


def smoke_test():
    print("[smoke test] 常物性纯导热，不接 D(C)")
    dr_cm = 0.005
    dt = 1.0
    t_end = 1800.0
    t_out, T_hist, r = solve_heat_constant(dr_cm=dr_cm, dt=dt, t_end=t_end)
    N = len(r) - 1
    print(f"网格：dr={dr_cm} cm  N={N} 单元，节点数 {N+1}，r∈[0,{R_CM}] cm")
    print(f"时间：dt={dt} s，共 {len(t_out)-1} 步，记录 {len(t_out)} 个时刻")
    print(f"热扩散系数 alpha = k/(rho*cp) = {alpha():.6e} m^2/s")
    for t_target in [0.0, 100.0, 300.0, 600.0, 900.0, 1200.0, 1500.0, 1800.0]:
        i = int(round(t_target / dt))
        T_center = T_hist[i, 0]
        T_surf = T_hist[i, N]
        T_air = T_air_at(t_target)
        print(f"  t={t_target:6.0f}s  中心 T={T_center:8.4f} °C  "
              f"表面 T={T_surf:8.4f} °C  烘房 T_air={T_air:8.4f} °C")
    T_final = T_hist[-1]
    ok_rise = T_final[0] > T_INIT
    ok_order = T_final[N] >= T_final[0]
    ok_bound = T_final.max() <= T_air_at(t_end) + 1e-9
    ok_phys = ok_rise and ok_order and ok_bound
    print(f"物理自检：中心升温={ok_rise}，表面≥中心={ok_order}，"
          f"不超烘房温度={ok_bound}  {'PASS' if ok_phys else 'FAIL'}")
    return ok_phys


if __name__ == "__main__":
    smoke_test()
