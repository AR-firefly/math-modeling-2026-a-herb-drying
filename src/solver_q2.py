# -*- coding: utf-8 -*-
"""
solver_q2.py —— 第二问正式求解器：变物性全时段（预热平衡 + 恒温干燥）热质耦合模型
================================================================================

实现依据：`docs/Q2_建模文档.md` v1.0（876 行）。**本文件不改动 `src/solver_q1.py`**
（该文件只读、禁改，用户拍板 Q2–Q4 保证 Q1 复现链完整）；Q1 的空间离散/测度/三对角
装配纪律在此**重写**（复制而非 import 修改），以便 Q2 独立演进。

----------------------------------------------------------------------------
一、控制方程（容量形式，建模文档 §3.1 (Q2-7) / §3.2 (Q2-8)）
----------------------------------------------------------------------------
    ρ(C)·c_p(C)·∂T/∂t = (1/r)·∂/∂r( r·k(C)·∂T/∂r )
    ∂C/∂t            = (1/r)·∂/∂r( r·D(C,T)·∂C/∂r )

物性（附录 3，逐位照抄，不做四舍五入）：
    ρ(C)   = 650 + 128C                             kg/m³
    c_p(C) = 1450 + 2736·C/(C+1)                    J/(kg·K)
    k(C)   = 0.21 + 0.38·C/(C+1)                    W/(m·K)
    D(C,T) = 2.4×10⁻³ · e^(−0.45/C) · e^(−3850/T)   m²/s

    ⚠️⚠️ **`D` 中的 `T` 必须用热力学温度 K（不是 ℃）**；输出阶段才转 ℃。
        （内部计算统一 SI；本文件所有 `+ 273.15` 处均带注释。）

边界（第三类，驱动量取 **n+1 隐式层**，Q1 §7.4 合并修正 M4）：
    −k(C_s)∂T/∂r|_R = h·(T_s − T_a(t))          h  = 25     W/(m²·K)   （附录 2）
    −D(C_s,T_s)∂C/∂r|_R = h_m·(C_s − C_a(t))    h_m= 8×10⁻⁷ m/s        （附录 2，口径 A）

    ⚠️ 附录 3 **未给** h / h_m；沿用附录 2 是「唯一可溯来源」的**推断**（Q2-A3），
       论文必须写成「沿用附录 2」而非「题给」。灵敏度扫描（§10.1）为必做项。

驱动量 `T_a(t)`、`C_a(t)`：附件 1（0–14400 s，步长 60 s）分段线性插值；
    `t > 14400 s` **冻结**在末值 `T_a = 50.165 ℃`、`C_a = 0.04986 kg/kg`（用户拍板决策 1）。

初值：`T = 28 ℃`、`C = 2.55 kg/kg` 均匀（题面逐位）；域 = 半径 2 cm 无限长圆柱，长 25 cm。
    **t = 0 起算，不承接 Q1 的 1800 s 终态**（Q2-A12：附录 2/3 物性不连续）。

----------------------------------------------------------------------------
二、数值方案（建模文档 §8 / §9）
----------------------------------------------------------------------------
空间：节点中心有限体积，节点 `r_i = iΔr`，`i = 0…N`；精确圆柱测度
      W_0 = Δr²/8、W_i = r_iΔr (1≤i≤N−1)、W_N = RΔr/2 − Δr²/8     (Q2-17)
时间：向后 Euler **全隐式**（M 矩阵 → 无条件稳定、极值原理无条件成立，无步长限制）
非线性：**外层块 Picard（先 T 后 C，块 Gauss–Seidel）**
      收敛判据（三项同时，Q2-23）：max|ΔT|、max|ΔC|、‖A φ − b‖_∞
      内迭代**恒冻结**：上一时间层 `T^n, C^n`（累积项右端）、几何、h/h_m、矩阵结构；
      可变的只有系数场 ω^{(m)}, k^{(m)}, D^{(m)}（含界面值）。
      **`D_i^(m) = D(C_i^(m), T_i^(m))`，其中 `T_i^(m)` 取 n+1 层的当前迭代值**
      （建模文档 §8.3-3a：写错就等于给 D 引入一阶滞后，与 M4 同族坑）。
      ω_i^(m) = ρ(C_i^(m))·c_p(C_i^(m)) —— 累积项的热容取**新时间层迭代值**；
      **上一时间层只出现在右端**（`T_i^n` / `C_i^n`）。

⚠️ 本实现对「取哪一层」的三条写死（建模文档 §15 交接清单第 3 条）：
   (1) 累积项右端恒取 **n 层**（`T^n` / `C^n`，Picard 全程不变）；
   (2) 边界驱动量恒取 **n+1 层**（`T_a^{n+1}` / `C_a^{n+1}`）；
   (3) 系数（含 `D` 的 `T`）取 **n+1 层当前迭代值**。
   本实现**二选一的细节**：Picard 第 m 次迭代内先解 T，再用**同一套系数场**解 C
   （即 `D` 用 `(C^{(m)}, T^{(m)})` 求值，不因刚更新的 `T^{(m+1)}` 重算）。
   理由：一次迭代只算一次指数；与 §8.3-3a 的字面口径一致。已写入 results 元数据。

`D` 的数值安全（§4.3）：
    · **对数空间求值**：`log D = log(2.4×10⁻³) − 0.45/C − 3850/T(K)`，再 `exp`。
      消除中间因子 `e^(−0.45/C)` 的过早下溢（直接乘积会先把仍可表示的结果错误地变成 0）。
    · **防卫下界** `D_floor = 1.0×10⁻²⁵ m²/s`，`D = exp(max(log D, log D_floor))`。
      非干扰论证：`D_floor / D(0.15, 323.15 K) = 1.25×10⁻¹⁶`，低于双精度相对分辨率。
    · **激活计数与留痕**（硬性）：任何 `log D < log D_floor` 都计数并留痕；
      **计数 > 0 即为异常信号**（Q2/Q3 交付区间内应为 0）→ 调用方必须停止并报告。

极值原理断言（§8.6 式 Q2-24，**越界即停、不许 clip**）：
    28 ℃ ≤ T ≤ max_t T_a(t)     0.01963 ≤ C ≤ 2.55 kg/kg

能量审计（§8.7 式 Q2-25，**不可沿用 Q1 的 ΣV_i(ρc_pT)_i 式——会报假违规**）：
    Σ_i W_i[ω_i^{n+1}T_i^{n+1} − ω_i^n T_i^n] − Σ_i W_i·T_i^n·[ω_i^{n+1} − ω_i^n]
        = Δt·(边界通量) + O(Δt²)
本实现逐内部步累计该残差。注意该式对**本离散格式是代数恒等式**（推导见模块 docstring
的 §8.7 复述）：把离散方程对 i 求和并用 `ω^{n+1}T^n = ω^nT^n + T^n(ω^{n+1}−ω^n)` 即得。
水分侧沿用 Q1 的 P3 守恒审计（跨界面通量逐对相消，**严格成立，不改**）。

----------------------------------------------------------------------------
三、时间步（建模文档 §9.3）
----------------------------------------------------------------------------
    I 段 `[0, 14400] s`：`Δt_I = 1/64 s`（921600 步）
    II 段 `(14400, t_end]`：`Δt_II = 1/8 s`（t_end = t_dry + 3 h）
    ⚠️ 表面节点局部响应 `0.368 s` **不是** Δt 约束（该节点被邻居导热以 169:1 强耦合，
       向后 Euler 在大 Δt 下给的是准稳态值）；**不得**用局部稳定性判据压 Δt（§7.2 注）。

终止（§11.4）：`C_max(t) ≤ 0.15` 首达后**再积分 3 h（10800 s）** 停止；
    硬上限 `t_cap = 14 d = 1 209 600 s`，到期未达标必须报「未达标」，**不得静默返回末行**。

----------------------------------------------------------------------------
四、`C_max(t)` 轨迹（Q3 衔接，§11.2）
----------------------------------------------------------------------------
**逐内部步**计算 `C_max, argmax, C(0,t), C(R,t), min C` 并落盘全程轨迹。
60 s 采样不合格；`Δt_II = 1/8 s` 给出 0.125 s 的括住区间，配线性插值定位 `t_dry`。

⚠️ `t_dry` 精度（§7.4）：`∂t_dry/∂C_max = 1/|dC_max/dt| = 9.536×10⁵ s per (kg/kg)`
   → 场值差 1×10⁻⁵ 就是 9.5 s；**`t_dry` 只能报 ±10 s 量级，Δt 再小也不改善**。

----------------------------------------------------------------------------
五、对比口径（§10.2，均为诊断，不进交付）
----------------------------------------------------------------------------
· `energy_form`：`capacity`（主，§3.1）/ `product_0C`、`product_K`（乘积形式两档 T_ref）
· `moisture_form`：`dC`（主，§3.2）/ `rhoC`（∂(ρC)/∂t 替代形式，量化 33.4% 速率差）
· `props_model`：`appendix3`（主）/ `appendix2`（退化回归 对比 1，复现 Q1 三锚点）
· `face_rule`：`arith`（主）/ `harm` / `faceCT`（Q2 新增：界面上同时取 C 与 T）

数据零近似：本文件所有附件/赛题数值逐位原样，不做四舍五入、不近似、不补题目未给参数。
"""
from __future__ import annotations

import math
import os
import sys
import time
from pathlib import Path

import numpy as np
from scipy.linalg import solve_banded
from scipy.linalg.lapack import dgtsv

if getattr(sys.stdout, "reconfigure", None):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]

# ============================================================================
# 0) 常量（附录 3 逐位；附录 2 的 h/h_m 为唯一可溯来源）
# ============================================================================
# ---- 附录 3 物性系数（逐位照抄）----
RHO_A, RHO_B = 650.0, 128.0
CP_A, CP_B = 1450.0, 2736.0
K_A, K_B = 0.21, 0.38
D0_AP3, BETA_AP3, E_D_AP3 = 2.4e-3, 0.45, 3850.0
LOG_D0_AP3 = math.log(D0_AP3)                 # = −6.032286541628237（与建模文档 §4.3 一致）

# ---- 附录 2（退化回归 对比 1 用；Q1 口径）----
A2_RHO, A2_CP, A2_K = 820.0, 2600.0, 0.36
A2_D0, A2_BETA = 7.0e-9, 0.89
LOG_A2_D0 = math.log(A2_D0)

# ---- 对流系数（附录 3 未给 → 沿用附录 2 = 唯一可溯来源；Q2-A3，属推断）----
H_DEFAULT = 25.0          # W/(m²·K)
HM_DEFAULT = 8.0e-7       # m/s

# ---- 几何与初值（题面逐位）----
R_CM = 2.0
R = R_CM / 100.0          # m = 0.02
LENGTH_CM = 25.0
T_INIT_C = 28.0           # ℃
C_INIT = 2.55             # kg/kg

# ---- 烘房驱动（附件 1 末行，逐位照抄；t > 14400 s 冻结）----
T_FREEZE_S = 14400.0
TA_FREEZE_C = 50.165
CA_FREEZE = 0.04986

# ---- 终止判据 ----
C_DRY = 0.15              # kg/kg（Q3 判据）
MARGIN_S = 10800.0        # 越过判据后再积分 3 h
T_CAP_S = 14.0 * 86400.0  # 1 209 600 s 硬上限

# ---- D 的防卫下界（§4.3）----
D_FLOOR = 1.0e-25
LOG_D_FLOOR = math.log(D_FLOOR)

# ---- 附件路径（可移植定位：MM_DATA_DIR 环境变量优先；否则在项目内/项目同级自动查找 A题 数据目录）----
def _find_data_dir() -> Path:
    if os.environ.get("MM_DATA_DIR"):
        return Path(os.environ["MM_DATA_DIR"])
    for _c in (ROOT / "A题", ROOT / "附件", ROOT.parent / "A题", Path.cwd() / "A题", Path.cwd()):
        if (_c / "附件" / "附件1.xlsx").is_file():
            return _c
    raise SystemExit(
        "未找到 A 题数据目录（须含 附件/附件1.xlsx）：请设置环境变量 MM_DATA_DIR 指向该目录，"
        "或把「A题」目录放到项目根/项目同级。"
    )


_DATA_DIR = str(_find_data_dir())
ATT1 = str(Path(_DATA_DIR) / "附件" / "附件1.xlsx")
ATT2 = str(Path(_DATA_DIR) / "附件" / "附件2.xlsx")
ATT3_DIR = Path(_DATA_DIR) / "附件" / "附件3"
RESULT2_TEMPLATE = str(ATT3_DIR / "result2.xlsx")
PDF = str(Path(_DATA_DIR) / "A题.pdf")

# ---- 附件 1 缓存 ----
_ATT1_CACHE = None
_DRIVER_CACHE = {}


def att1():
    """返回 (t, T_a, C_a)：附件 1 原始采样（241 行，t = 0, 60, …, 14400 s）。"""
    global _ATT1_CACHE
    if _ATT1_CACHE is None:
        import openpyxl
        wb = openpyxl.load_workbook(ATT1, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        t = np.array([float(r[0]) for r in rows])
        T = np.array([float(r[1]) for r in rows])
        C = np.array([float(r[2]) for r in rows])
        _ATT1_CACHE = (t, T, C)
    return _ATT1_CACHE


def driver_table():
    """段内插值查表（为速度写成纯标量算术，不调 np.interp）。"""
    global _DRIVER_CACHE
    if not _DRIVER_CACHE:
        t, T, C = att1()
        dt_s = float(t[1] - t[0])
        nseg = int(round(t[-1] / dt_s))            # 240 段
        if abs(t[0]) > 1e-12 or abs(t[-1] - T_FREEZE_S) > 1e-9:
            raise ValueError("附件 1 时间轴不是 0…14400 s（步长 %g）" % dt_s)
        _DRIVER_CACHE = dict(dt_s=dt_s, nseg=nseg, T=T.copy(), C=C.copy(),
                             inv_dt=1.0 / dt_s, Ta_max=float(T.max()),
                             Ca_min=float(C.min()), Ca_max=float(C.max()))
    return _DRIVER_CACHE


def driver_at(t, interp="linear"):
    """烘房驱动 (T_a [℃], C_a [kg/kg])，标量；分段线性 + 14400 s 后冻结末值。

    `interp='step'` 为**误差分析用的对照**：零阶保持（取左侧采样值），
    用于量化「附件 1 的读取方式」这项**数据误差**（§10 第 4 类）。
    主口径恒为 'linear'（分段线性，与 Q1 §6.1 一致）。
    """
    if t >= T_FREEZE_S:
        return TA_FREEZE_C, CA_FREEZE
    d = driver_table()
    x = t * d["inv_dt"]
    i = int(x)
    if i >= d["nseg"]:                              # 浮点边界保护
        i = d["nseg"] - 1
    if interp == "step":                            # 零阶保持（对照）
        return float(d["T"][i]), float(d["C"][i])
    f = x - i
    T = float(d["T"][i] + (d["T"][i + 1] - d["T"][i]) * f)
    C = float(d["C"][i] + (d["C"][i + 1] - d["C"][i]) * f)
    return T, C


# ============================================================================
# 1) 物性场（附录 3 主口径 / 附录 2 退化口径）
# ============================================================================
def default_params():
    """锁定参数；灵敏度扫描通过 `**overrides` 覆盖（**不改动默认值**）。"""
    return dict(
        # 附录 3 物性系数
        RHO_A=RHO_A, RHO_B=RHO_B, CP_A=CP_A, CP_B=CP_B, K_A=K_A, K_B=K_B,
        D0=D0_AP3, BETA=BETA_AP3, E_D=E_D_AP3,
        # log 形式在 solve_q2 里由 P["D0"] 现算（保证 D0 可被 **overrides 覆盖）
        LOG_D0=LOG_D0_AP3, LOG_A2_D0=LOG_A2_D0,
        # 附录 2 退化口径参数（仅 props_model='appendix2' 用）
        A2_RHO=A2_RHO, A2_CP=A2_CP, A2_K=A2_K, A2_D0=A2_D0, A2_BETA=A2_BETA,
        # 对流系数
        H=H_DEFAULT, HM=HM_DEFAULT,
        # 口径选择
        props_model="appendix3",       # 'appendix3' | 'appendix2'
        energy_form="capacity",        # 'capacity' | 'product_0C' | 'product_K'
        moisture_form="dC",            # 'dC' | 'rhoC'
        face_rule="arith",             # 'arith' | 'harm' | 'faceCT'
        # 驱动情景（对比 6 / §10.1）
        ta_scenario="freeze_last",     # 'freeze_last' | 'tail_mean' | 'row_min' | 'row_max' | 'const'
        ca_scenario="freeze_last",
        const_ta=TA_FREEZE_C, const_ca=CA_FREEZE,
        # 灵敏度专用（默认 1.0；见 omega_of_C 注释）
        RHO_SCALE=1.0, CP_SCALE=1.0,
        # 附件 1 的读取方式（'linear' 主口径；'step' 只作数据误差对照）
        INTERP="linear",
        # 初值
        T_INIT=T_INIT_C, C_INIT=C_INIT,
    )


def rho_of_C(C, P):
    if P["props_model"] == "appendix2":
        return np.full_like(C, P["A2_RHO"])
    return P["RHO_A"] + P["RHO_B"] * C


def omega_of_C(C, P):
    """有效体积热容 ω(C) = ρ(C)·c_p(C) [J/(m³·K)]。

    `RHO_SCALE` / `CP_SCALE` 只用于灵敏度扫描（§10.1「ρ 与 c_p 的系数各 ±10 %」），
    默认为 1.0，**不改变主口径**。ρ 与 c_p 在 Q2 中只通过 ω 进入方程（k 不含 ρ、`dC`
    形式的扩散方程不含 ρ），故两个扫描都是「能量时间常数」的 ±10 % 扫描。
    """
    if P["props_model"] == "appendix2":
        base = np.full_like(C, P["A2_RHO"] * P["A2_CP"])
    else:
        base = (P["RHO_A"] + P["RHO_B"] * C) * (P["CP_A"] + P["CP_B"] * C / (C + 1.0))
    sc = P.get("RHO_SCALE", 1.0) * P.get("CP_SCALE", 1.0)
    return base if sc == 1.0 else base * sc


def k_of_C(C, P):
    """导热系数 k(C) [W/(m·K)]。"""
    if P["props_model"] == "appendix2":
        return np.full_like(C, P["A2_K"])
    return P["K_A"] + P["K_B"] * C / (C + 1.0)


def logD_of_CT(C, T_C, P):
    """log D（**对数空间求值**，§4.3）。

    ⚠️ `T_C` 为摄氏温度，本函数内部 `+273.15` 转 K 后再进 `e^(−3850/T)`——**必须用 K**。
    """
    if P["props_model"] == "appendix2":
        # 附录 2：D = 7×10⁻⁹·e^(−0.89/C)，无 T 依赖（退化回归口径）
        return P["LOG_A2_D0"] - P["A2_BETA"] / C
    T_K = T_C + 273.15                              # ← ℃ → K（D 的指数项要求 K）
    # ⚠️ 必须用 P["LOG_D0"]（= log(P["D0"])）而**不是**模块常量 LOG_D0_AP3 ——
    #    否则 D₀ 的灵敏度扫描会被**静默忽略**（2026-09-11 实测踩到：D0×0.5 与 ×2
    #    给出逐位相同的结果）。默认时两者数值相同，故不影响交付值。
    return P["LOG_D0"] - P["BETA"] / C - P["E_D"] / T_K


class FloorMonitor:
    """`D_floor` 激活留痕（§4.3）。激活次数 > 0 即异常 → 调用方必须停止并报告。"""

    def __init__(self):
        self.count = 0
        self.events = []
        self.max_events = 20

    def note(self, mask, logD, C, T_C, t):
        n = int(np.count_nonzero(mask))
        if n == 0:
            return
        self.count += n
        for i in np.flatnonzero(mask):
            if len(self.events) >= self.max_events:
                break
            self.events.append(dict(t_s=float(t), i=int(i), C=float(C[i]),
                                    T_C=float(T_C[i]), logD=float(logD[i]),
                                    logD_floor=LOG_D_FLOOR))

    def to_dict(self):
        return dict(count=int(self.count), log_D_floor=LOG_D_FLOOR, D_floor=D_FLOOR,
                    note=("激活次数必须为 0（建模文档 §4.3：极值原理给 C ≥ 0.01963，"
                          "是下溢点的 31.7 倍）。> 0 表示前提被推翻 → 停止并回报。"),
                    first_events=self.events)


def eval_D(logD, monitor, C, T_C, t):
    """把 logD 落到下界之上并留痕，返回 D 数组。"""
    if float(logD.min()) < LOG_D_FLOOR:
        monitor.note(logD < LOG_D_FLOOR, logD, C, T_C, t)
        logD = np.maximum(logD, LOG_D_FLOOR)
    return np.exp(logD)


def interface_gam(node, C, T_C, rule, kind, P):
    """界面系数（§8.5）：`arith` / `harm` / `faceCT`。

    faceCT（**Q2 新增**）：在界面上同时取 C 与 T 再求 k 或 D
        —— Q1 的 faceC 只对 C 取界面值；Q2 的 D 显含 T，故必须同时给界面的 T。
    """
    if rule == "arith":
        return 0.5 * (node[:-1] + node[1:])
    if rule == "harm":
        a, b = node[:-1], node[1:]
        return 2.0 * a * b / (a + b)
    if rule == "faceCT":
        Cf = 0.5 * (C[:-1] + C[1:])
        Tf = 0.5 * (T_C[:-1] + T_C[1:])
        if kind == "k":
            return k_of_C(Cf, P)
        return np.exp(np.maximum(logD_of_CT(Cf, Tf, P), LOG_D_FLOOR))
    raise ValueError("未知界面系数规则 face_rule=%r" % (rule,))


def dimensionless():
    """复核建模文档 §7.1 的无量纲数与 §7.2 的特征时间（用于与文档互检，进 results 元数据）。

    `α = k/(ρc_p)`、`Bi_h = hR/k`、`Bi_m = h_mR/D`、`Le = α/D`。
    ⚠️ `D` 求值处的 `T` 用 **K**（本函数内部 `+273.15`）。
    """
    P = default_params()
    rows = []
    for C_val, T_C in ((2.55, 28.0), (2.55, 50.0), (1.00, 50.0), (0.15, 50.0)):
        Cv = np.array([C_val])
        rho = float(rho_of_C(Cv, P)[0])
        cp = float((P["CP_A"] + P["CP_B"] * Cv / (Cv + 1.0))[0])
        kv = float(k_of_C(Cv, P)[0])
        Dv = float(np.exp(logD_of_CT(Cv, np.array([T_C]), P)[0]))
        a = kv / (rho * cp)
        dh = math.degrees(0)  # noqa: F841  (占位，保持可读)
        rows.append(dict(C=C_val, T_C=T_C, rho=rho, cp=cp, k=kv, D=Dv,
                         alpha=a, Bi_h=P["H"] * R / kv, Bi_m=P["HM"] * R / Dv,
                         Le=a / Dv,
                         tau_heat_s=R * R / a, tau_mass_s=R * R / Dv))
    extra = dict(
        R_m=R, h=P["H"], h_m=P["HM"],
        tau_mass_at_C015_s=R * R / rows[3]["D"],
        dCmax_dt_at_015=1.0471756754118161e-05 * (0.15 - 0.04986),
        dt_dry_per_dCmax_s=1.0 / (1.0471756754118161e-05 * (0.15 - 0.04986)),
        note="对照 docs/Q2_建模文档.md §7.1/§7.2/§7.4 的数值；差异即为文档取整或口径差",
    )
    return dict(rows=rows, extra=extra)


# ============================================================================
# 2) 网格与测度（Q2-17）
# ============================================================================
def build_grid(dr_cm):
    """均匀径向网格 + 精确圆柱测度 W_i = ∫_{r_{i−1/2}}^{r_{i+1/2}} r dr。

    中心测度必须是 Δr²/8（不得写 Δr²/4）；表面测度必须是 RΔr/2 − Δr²/8
    （不得粗网格近似成 RΔr/2）——Q1 §7.5 的两条纪律。
    """
    dr = dr_cm / 100.0
    N = int(round(R / dr))
    if abs(N * dr - R) > 1e-12:
        raise ValueError("dr_cm=%g 不能整除 R=%g cm" % (dr_cm, R_CM))
    r = np.arange(N + 1) * dr
    W = np.empty(N + 1)
    W[0] = dr * dr / 8.0
    W[1:N] = r[1:N] * dr
    W[N] = R * dr / 2.0 - dr * dr / 8.0
    return r, dr, N, W


def node_index_for(r_cm, dr_cm):
    """`r_cm` 对应的节点下标（要求落在节点上，**不插值**）。"""
    q = r_cm / dr_cm
    i = int(round(q))
    if abs(q - i) > 1e-9:
        raise ValueError("r=%g cm 不在 dr=%g cm 的节点上" % (r_cm, dr_cm))
    return i


# ============================================================================
# 3) 三对角装配（温度/水分共用；见模块 docstring 的离散式）
# ============================================================================
def assemble(cap, gam_face, h_coef, dt, geo, form="capacity", omega_old=None):
    """装配三对角带矩阵。

    方程写为：`d_i φ_i^{n+1} + u_i φ_{i+1}^{n+1} + l_i φ_{i−1}^{n+1} = rhs_i`
    `rhs_i = w_i·φ_i^n`（`w_i` = 旧层权重 / 新层权重；capacity 形式 `w_i ≡ 1`）。
    表面行另有驱动项 `+ beta·φ_driver`。

    参数
    ----------
    cap      : (N+1,) 新层容量（能量 = ω^{(m)}；水分 dC 形式 = 1；水分 rhoC 形式 = ρ^{(m)}）
    gam_face : (N,)   界面输运系数（能量 = k_{i+1/2}；水分 = D_{i+1/2}）
    geo      : dict（由 `make_geo` 预计算，与 dt 无关）
    form     : 'capacity' | 'product'（product 时右端权重 w_i = ω_i^n/ω_i^{(m)}）

    返回 (ab, beta, d, l, u, w)
    """
    N = geo["N"]
    G1, G2, ax4, sfac, sbnd, W = geo["G1"], geo["G2"], geo["ax4"], geo["sfac"], geo["sbnd"], geo["W"]
    inv_cap = 1.0 / cap

    d = np.empty(N + 1)
    l = np.zeros(N + 1)
    u = np.zeros(N + 1)
    w = np.ones(N + 1)

    # ---- 内部节点 i = 1..N−1 ----
    a1 = dt * G1 * gam_face[1:N] * inv_cap[1:N]        # Γ_{i+1/2}
    a2 = dt * G2 * gam_face[0:N - 1] * inv_cap[1:N]    # Γ_{i−1/2}
    d[1:N] = 1.0 + a1 + a2
    u[1:N] = -a1
    l[1:N] = -a2
    if form == "product" and omega_old is not None:
        w[1:N] = omega_old[1:N] * inv_cap[1:N]

    # ---- 轴节点 i = 0（半控制体 [0, Δr/2]，对称面通量为零）----
    a0 = dt * ax4 * gam_face[0] * inv_cap[0]
    d[0] = 1.0 + a0
    u[0] = -a0
    if form == "product" and omega_old is not None:
        w[0] = omega_old[0] * inv_cap[0]

    # ---- 表面节点 i = N（半控制体 [R−Δr/2, R]，Robin）----
    alpha = dt * sfac * gam_face[N - 1] * inv_cap[N]
    beta = dt * sbnd * h_coef * inv_cap[N]
    d[N] = 1.0 + alpha + beta
    l[N] = -alpha
    if form == "product" and omega_old is not None:
        w[N] = omega_old[N] * inv_cap[N]

    ab = np.zeros((3, N + 1))
    ab[0, 1:] = u[:-1]          # u[i]（行 i 中 φ_{i+1} 的系数）→ ab[0, i+1]
    ab[1, :] = d
    ab[2, :-1] = l[1:]          # l[i]（行 i 中 φ_{i−1} 的系数）→ ab[2, i−1]
    return ab, beta, d, l, u, w


def make_geo(r, dr, N, W):
    """与 dt 无关的几何因子（预计算一次）。"""
    rr = r[1:N]
    return dict(N=N, W=W,
                G1=(rr + dr / 2.0) / (rr * dr * dr),      # r_{i+1/2}/(r_i Δr²)
                G2=(rr - dr / 2.0) / (rr * dr * dr),      # r_{i−1/2}/(r_i Δr²)
                ax4=4.0 / (dr * dr),                      # 轴节点：4Γ_{1/2}/(Δr² cap_0)
                sfac=(R - dr / 2.0) / (dr * W[N]),        # 表面内侧：Γ_{N−1/2}/(Δr W_N)
                sbnd=R / W[N])                            # 表面外侧：h/(W_N)


def tdsolve(d, l, u, rhs):
    """解三对角系统 `d φ + u·φ^{+1} + l·φ^{−1} = rhs`（等价于 `solve_banded((1,1),…)`）。

    ⚠️ 用 `scipy.linalg.lapack.dgtsv` 而非 `solve_banded`：两者结果**逐位相同**
    （本机实测 max|Δ| = 0.0），但 dgtsv 3.7 µs vs solve_banded 17.9 µs（N=201）。
    生产全程调用次数 = 6 × 步数 量级，该差异直接决定全流程墙钟时间。
    `rhs` 必须是形状 (N+1,) 的连续数组；返回其内存上的视图（调用方不再复用 rhs）。
    """
    n = rhs.size
    x = dgtsv(l[1:], d, u[:-1], rhs.reshape(n, 1), overwrite_b=1)[3]
    return x[:, 0]


def banded_residual(d, l, u, phi, rhs):
    """‖A φ − b‖_∞（收敛判据第 3 项，Q2-23）。"""
    res = d * phi
    res[:-1] += u[:-1] * phi[1:]
    res[1:] += l[1:] * phi[:-1]
    res -= rhs
    return float(np.abs(res).max())


# ============================================================================
# 4) 主求解器
# ============================================================================
class _Grow:
    """可增长一维 float64 数组（逐内部步轨迹；避免 list 的 ~32 B/元素开销）。"""

    def __init__(self, cap=4096):
        self.a = np.empty(cap)
        self.n = 0

    def append(self, v):
        if self.n == self.a.size:
            self.a = np.resize(self.a, self.a.size * 2)
        self.a[self.n] = v
        self.n += 1

    def array(self):
        return self.a[:self.n].copy()


def scenario_driver(scen, samples, const):
    """附件 1 外推情景（对比 6 / §10.1 的 T_a、C_a 平台档）。

    'freeze_last'         : 主口径——t > 14400 s 冻结末值（样本不动，与主求解器同源）
    'tail_mean'           : E1——末段（最后 1 h = 60 个采样）均值
    'row_min'/'row_max'   : E2——末段最小/最大值
    'const'               : 全程取常数 const
    """
    s = np.asarray(samples, float)
    if scen == "freeze_last":
        return s.copy()
    if scen == "const":
        return np.full_like(s, float(const))
    tail = s[-60:]
    if scen == "tail_mean":
        return np.concatenate([s[:-60], np.full(60, float(tail.mean()))])
    if scen == "row_min":
        return np.concatenate([s[:-60], np.full(60, float(tail.min()))])
    if scen == "row_max":
        return np.concatenate([s[:-60], np.full(60, float(tail.max()))])
    raise ValueError("未知情景 %r" % (scen,))


def solve_q2(dr_cm=0.01, dt1=1.0 / 64.0, t_switch=T_FREEZE_S, dt2=1.0 / 8.0,
             t_cap=T_CAP_S, margin_s=MARGIN_S, c_dry=C_DRY,
             stop_on_dry=True, t_end_override=None, t_start_state=None,
             picard_tol=1e-12, picard_max=100,
             snap_every_s=1.0, snap_until_s=10800.0, snap_end=None, snap_nodes=None,
             profile_times=(), cmax_track=True, checkpoints=(), audit=True, audit_every=1,
             energy_form=None, moisture_form=None, face_rule=None,
             verbose=True, progress_every=0, **overrides):
    """问题 2 主求解器（容量形式 + 块 Picard + 两段式时间步）。

    参数
    ----------
    dr_cm          : 空间步长 [cm]，须整除 0.1 cm；主网格 0.01（N=200，输出点全落节点）
    dt1, dt2       : I 段 / II 段内部时间步 [s]
    t_switch       : 两段切换时刻 [s]（默认 14400 = 附件 1 数据终点）
    t_cap          : 硬上限 [s]（默认 14 d）；到时未达标 → status='not_dry_by_cap'
    margin_s       : 达标后再积分的余量 [s]（默认 3 h）
    stop_on_dry    : True = 主生产（监测 C_max 并提前停止）；False = 跑满 `t_end_override`
    t_start_state  : dict(t, T, C, cum_flux, mass, n_steps) 重启（检查点续算）
    snap_every_s   : 场快照间隔 [s]（0 = 不存）
    snap_until_s   : 场快照截止时刻 [s]
    snap_nodes     : 只存这些节点下标（None = 全部 N+1 个）。用于网格无关性研究的
                     公共交付半径（21 个），避免在 N=800 时存下 10801×801 的大数组
    profile_times  : 在这些时刻额外存**全场**剖面（键 profiles = {t: (T_full, C_full)}）
    snap_end       : 若给定，则快照覆盖到 t_end_override（单段基准用）
    cmax_track     : 是否逐内部步记录 C_max 轨迹（Q3 需要）
    checkpoints    : 需要落检查点的时刻 tuple [s]
    audit          : 是否做能量/水分审计（每 `audit_every` 步）

    返回 dict（键名见文件末 `KEYS` 与 run_q2.py）
    """
    t_wall0 = time.time()
    P = default_params()
    P.update(overrides)
    P["LOG_D0"] = math.log(P["D0"])                 # D₀ 可被 overrides 覆盖 → log 必须现算
    P["LOG_A2_D0"] = math.log(P["A2_D0"])
    if energy_form is not None:
        P["energy_form"] = energy_form
    if moisture_form is not None:
        P["moisture_form"] = moisture_form
    if face_rule is not None:
        P["face_rule"] = face_rule

    r, dr, N, W = build_grid(dr_cm)
    geo = make_geo(r, dr, N, W)

    # ⚠️ 参数白名单校验：**禁止静默回退**。曾因 `moisture_form="rhoC_full"` 这类
    #    拼写错误被静默当成 capacity，跑出「看似正常但其实是别的口径」的结果
    #    （2026-09-11 实测踩到）。未知取值一律报错。
    if P["energy_form"] not in ("capacity", "product_0C", "product_K"):
        raise ValueError("未知 energy_form=%r" % (P["energy_form"],))
    if P["moisture_form"] not in ("dC", "rhoC"):
        raise ValueError("未知 moisture_form=%r" % (P["moisture_form"],))
    if P["face_rule"] not in ("arith", "harm", "faceCT"):
        raise ValueError("未知 face_rule=%r" % (P["face_rule"],))
    if P["props_model"] not in ("appendix3", "appendix2"):
        raise ValueError("未知 props_model=%r" % (P["props_model"],))
    if P.get("INTERP", "linear") not in ("linear", "step"):
        raise ValueError("未知 INTERP=%r" % (P.get("INTERP"),))
    T_ref = 0.0 if P["energy_form"] != "product_K" else 273.15
    e_form = "product" if P["energy_form"].startswith("product") else "capacity"
    m_form = "product" if P["moisture_form"] == "rhoC" else "capacity"

    # ---- 驱动（含外推情景）----
    drv = driver_table()
    Ta_scen = scenario_driver(P["ta_scenario"], drv["T"], P["const_ta"])
    Ca_scen = scenario_driver(P["ca_scenario"], drv["C"], P["const_ca"])
    use_frozen_tbl = (P["ta_scenario"] == "freeze_last" and P["ca_scenario"] == "freeze_last")
    _inv_dt, _nseg = drv["inv_dt"], drv["nseg"]

    interp_kind = P.get("INTERP", "linear")

    def drv_at(t):
        if use_frozen_tbl:
            return driver_at(t, interp_kind)
        if t >= T_FREEZE_S:
            return float(Ta_scen[-1]), float(Ca_scen[-1])
        x = t * _inv_dt
        i = int(x)
        if i >= _nseg:
            i = _nseg - 1
        f = x - i
        return (float(Ta_scen[i] + (Ta_scen[i + 1] - Ta_scen[i]) * f),
                float(Ca_scen[i] + (Ca_scen[i + 1] - Ca_scen[i]) * f))

    # ---- 极值原理界（Q2-24，由驱动样本现算，不写死）----
    ta_bound = max(float(np.max(Ta_scen)), float(Ta_scen[-1]))
    ca_bound = min(float(np.min(Ca_scen)), float(Ca_scen[-1]))
    c_lo = min(float(P["C_INIT"]), ca_bound)
    t_lo = min(float(P["T_INIT"]), float(np.min(Ta_scen)))
    t_hi = max(float(P["T_INIT"]), ta_bound)

    # ---- 状态 ----
    t = 0.0
    step = 0
    cum_flux = 0.0
    if t_start_state is not None:
        t = float(t_start_state["t"])
        step = int(t_start_state.get("n_steps", 0))
        T = np.array(t_start_state["T"], float)
        C = np.array(t_start_state["C"], float)
        cum_flux = float(t_start_state.get("cum_flux", 0.0))
        if T.size != N + 1:
            raise ValueError("检查点网格与 dr_cm=%g 不符" % dr_cm)
    else:
        T = np.full(N + 1, float(P["T_INIT"]))
        C = np.full(N + 1, float(P["C_INIT"]))
    Tt = T - T_ref
    om_n = omega_of_C(C, P)
    m0_ref = float(np.dot(W, C))       # 守恒审计基准（重启时即检查点质量，保持 P3 自洽）

    floor_mon = FloorMonitor()
    viol = []

    # ---- 快照 ----
    # 快照按**物理时刻**调度（不是按步数！）——`Δt` 在段切换 / 检查点重启后会变，
    # 若用 `step % snap_gap` 调度，重启后的采样间隔会随 Δt 缩放（2026-09-11 实测踩到：
    # `dtII_*` 自 t=14400 重启后，Δt 从 1/64 变 1/8 → 采样间隔 60 s 变成 480 s）。
    snap_stop = snap_until_s if snap_end is None else snap_end
    sn = np.arange(N + 1) if snap_nodes is None else np.asarray(snap_nodes, int)
    snap_t, snap_T, snap_C = [], [], []
    snap_next = None
    if snap_every_s > 0:
        snap_next = (math.floor(t / snap_every_s + 1e-12) + 1) * snap_every_s
        if t == 0.0 and snap_stop >= 0.0:
            snap_t.append(0.0)
            snap_T.append(T[sn].copy())
            snap_C.append(C[sn].copy())

    # ---- 逐内部步轨迹 ----
    tr = {k: (_Grow() if cmax_track else None)
          for k in ("t", "cmax", "arg", "c0", "cR", "cmin")}

    ckpt = {}
    profiles = {}
    prof_left = sorted(float(v) for v in profile_times)
    picard_iters = _Grow()
    n_not_conv = 0
    res_max_T = res_max_C = 0.0
    dT_seen = dC_seen = 0.0

    # ---- 审计 ----
    E_prev = float(np.dot(W, om_n * Tt))
    aud = dict(n=0, max_abs=0.0, max_rel_perstep=0.0, sum_abs=0.0, samples=[],
               peak_dE=0.0, peak_flux=0.0)

    # ---- 时间推进 ----
    dt_now = dt1
    n_switch = 0
    t_end = None
    t_dry = None
    t_dry_prev = None
    cmax_prev = None
    status = "unknown"
    n_steps_total = step
    mass = float(np.dot(W, C))

    while True:
        # ---------- 终止判定（推进之前）----------
        if stop_on_dry:
            if t_dry is not None and t >= t_end - 1e-9:
                status = "ok"
                break
        else:
            if t_end_override is not None and t >= t_end_override - 1e-9:
                status = "fixed_end"
                break
        if t >= t_cap - 1e-9:
            status = "not_dry_by_cap"
            break
        # ---------- 段切换 ----------
        if n_switch == 0 and dt2 != dt1 and t >= t_switch - 1e-12:
            dt_now = dt2
            n_switch = 1

        t_new = t + dt_now
        ta_new, ca_new = drv_at(t_new)

        # ---------------- 块 Picard（先 T 后 C）----------------
        k_it = 0
        converged = False
        T_it, C_it = Tt, C            # 初值 = 上一时间层（§8.3-2）
        T_prev, C_prev = T_it, C_it
        ones = np.ones(N + 1)
        T_old_t = T - T_ref           # ⛔ 右端的 n 层（T̃ 坐标），Picard 全程不变
        for m in range(1, picard_max + 1):
            k_it = m
            # (a) 系数场：ω^{(m)}, k^{(m)}, D^{(m)}（D 的 T 取 n+1 层**当前迭代值**）
            om = omega_of_C(C_it, P)
            kv = k_of_C(C_it, P)
            Dv = eval_D(logD_of_CT(C_it, T_it + T_ref, P), floor_mon,
                        C_it, T_it + T_ref, t_new)
            # (b) 温度
            k_face = interface_gam(kv, C_it, T_it + T_ref, P["face_rule"], "k", P)
            ab, beta, d, l, u, w = assemble(
                om, k_face, P["H"], dt_now, geo, e_form,
                om_n if e_form == "product" else None)
            rhs = w * T_old_t                      # ⛔ 右端恒取 n 层
            rhs[N] += beta * (ta_new - T_ref)
            T_it = tdsolve(d, l, u, rhs)
            # (c) 水分（**同一套系数场**：D 用 (C^{(m)}, T^{(m)})，见模块 docstring 声明）
            if m_form == "product":
                cap_m = rho_of_C(C_it, P)
                om_old_m = rho_of_C(C, P)          # n 层 ρ
            else:
                cap_m = ones
                om_old_m = None
            D_face = interface_gam(Dv, C_it, T_it + T_ref, P["face_rule"], "D", P)
            if m_form == "product":
                # ∂(ρC)/∂t = ∇·(ρD∇C)：界面输运系数必须是 **ρ_face·D_face**。
                # 只把 D 当系数会退化成 ∇·(D∇C)（第三个模型，不是 §3.2 的替代形式）。
                # 2026-09-11 实测：缺 ρ_face 时表面 C(1800 s) = 2.520 而非 1.649（几乎不干燥）。
                rho_it = rho_of_C(C_it, P)
                D_face = D_face * 0.5 * (rho_it[:-1] + rho_it[1:])
            ab2, beta2, d2, l2, u2, w2 = assemble(
                cap_m, D_face, (P["HM"] * rho_of_C(C_it, P)[N]) if m_form == "product"
                else P["HM"], dt_now, geo, m_form, om_old_m)
            rhs2 = w2 * C                          # ⛔ 右端恒取 n 层
            rhs2[N] += beta2 * ca_new
            C_it = tdsolve(d2, l2, u2, rhs2)

            dT = float(np.abs(T_it - T_prev).max())
            dC = float(np.abs(C_it - C_prev).max())
            T_prev, C_prev = T_it, C_it
            dT_seen = max(dT_seen, dT)
            dC_seen = max(dC_seen, dC)
            if dT < picard_tol and dC < picard_tol:
                converged = True
                break

        # 残差（判据第 3 项）：在最终接受的迭代值上算一次
        om = omega_of_C(C_it, P)
        kv = k_of_C(C_it, P)
        Dv = eval_D(logD_of_CT(C_it, T_it + T_ref, P), floor_mon, C_it, T_it + T_ref, t_new)
        k_face = interface_gam(kv, C_it, T_it + T_ref, P["face_rule"], "k", P)
        ab, beta, d, l, u, w = assemble(om, k_face, P["H"], dt_now, geo, e_form,
                                        om_n if e_form == "product" else None)
        rhs = w * T_old_t
        rhs[N] += beta * (ta_new - T_ref)
        res_max_T = max(res_max_T, banded_residual(d, l, u, T_it, rhs))
        if m_form == "product":
            cap_m_use = rho_of_C(C_it, P)
            om_old_m = rho_of_C(C, P)
        else:
            cap_m_use = ones
            om_old_m = None
        D_face = interface_gam(Dv, C_it, T_it + T_ref, P["face_rule"], "D", P)
        if m_form == "product":                       # 同上一处：界面系数必须含 ρ_face
            rho_it = rho_of_C(C_it, P)
            D_face = D_face * 0.5 * (rho_it[:-1] + rho_it[1:])
        ab2, beta2, d2, l2, u2, w2 = assemble(
            cap_m_use, D_face,
            (P["HM"] * rho_of_C(C_it, P)[N]) if m_form == "product" else P["HM"],
            dt_now, geo, m_form, om_old_m)
        rhs2 = w2 * C
        rhs2[N] += beta2 * ca_new
        res_max_C = max(res_max_C, banded_residual(d2, l2, u2, C_it, rhs2))

        picard_iters.append(k_it)
        if not converged:
            n_not_conv += 1

        # ---------------- 接受 n+1 层 ----------------
        T_new = T_it + T_ref
        C_new = C_it
        cmax_now = float(C_new.max())
        cmin_now = float(C_new.min())
        tmin_now = float(T_new.min())
        tmax_now = float(T_new.max())

        # ---------------- 极值原理断言（越界即停，不许 clip）----------------
        # ⚠️ 该断言的**证明前提**是 §8.6 的 M 矩阵性质，只对**交付口径**
        #    （能量 = 容量形式、水分 = ∂C/∂t）成立。乘积形式 `∂(ωT)/∂t = ∇·(k∇T)`
        #    的右端权重是 ω^n/ω^{n+1}（可 >1），**不是**凸组合 → M 矩阵证明不适用，
        #    该形式「越界」是它自身的形式缺陷（正是 §3.1.2 判它不可用的理由之一），
        #    故只在交付口径下强制断言（2026-09-11：曾因对 product 形式也断言而误报）。
        enforce_extremum = (e_form == "capacity" and m_form == "capacity")
        if enforce_extremum and (cmin_now < c_lo - 1e-12
                                 or cmax_now > float(P["C_INIT"]) + 1e-12):
            viol.append(dict(t_s=float(t_new), kind="C", min=cmin_now, max=cmax_now,
                             bound=[c_lo, float(P["C_INIT"])]))
        if enforce_extremum and (tmin_now < t_lo - 1e-9 or tmax_now > t_hi + 1e-9):
            viol.append(dict(t_s=float(t_new), kind="T", min=tmin_now, max=tmax_now,
                             bound=[t_lo, t_hi]))
        if viol:
            status = "extremum_violation"
            t, n_steps_total = t_new, step + 1
            break

        # ---------------- 审计（Q2-25 + Q1 的 P3）----------------
        om_new = omega_of_C(C_new, P)
        flux_step = dt_now * R * P["H"] * (ta_new - float(T_new[N]))     # 边界导热净流入
        E_new = float(np.dot(W, om_new * (T_new - T_ref)))
        if audit and (step % audit_every == 0):
            cap_acct = float(np.dot(W, T * (om_new - om_n)))             # 「热容变化在旧温度下记的账」
            dE = E_new - E_prev
            resid = dE - cap_acct - flux_step
            scale = max(abs(dE), abs(cap_acct), abs(flux_step), 1e-300)
            aud["n"] += 1
            aud["max_abs"] = max(aud["max_abs"], abs(resid))
            aud["max_rel_perstep"] = max(aud["max_rel_perstep"], abs(resid) / scale)
            aud["peak_dE"] = max(aud["peak_dE"], abs(dE))
            aud["peak_flux"] = max(aud["peak_flux"], abs(flux_step))
            aud["sum_abs"] += abs(resid)
            if len(aud["samples"]) < 5:
                aud["samples"].append(dict(t_s=float(t_new), dE=dE, cap_account=cap_acct,
                                           boundary_flux=flux_step, residual=resid))
        E_prev = E_new

        # 水分侧守恒（跨界面通量逐对相消，严格成立；沿用 Q1 的 P3 恒等式，不改）
        # 本文件的 W_i 是**真实测度** ∫r dr（Q2-17），故恒等式为
        #     Σ_i W_i (C_i^n − C_i^{n+1}) = Δt·R·q_surf ,   q_surf = h_m(C_N − C_a) > 0
        # （Q1 的 TVOL = 2W 且通量用 2R q，两边同乘 2 → 相对误差不受影响；本文件不沿用该因子。）
        q_surf = P["HM"] * (float(C_new[N]) - ca_new)        # 表面流失通量（C 单位，m/s）
        cum_flux += dt_now * R * q_surf
        mass = float(np.dot(W, C_new))

        T, C, om_n = T_new, C_new, om_new
        Tt = T - T_ref
        mass_prev = cmax_prev
        t, step = t_new, step + 1
        n_steps_total = step

        # ---------------- C_max 轨迹（逐内部步，§11.2）----------------
        if cmax_track:
            tr["t"].append(t); tr["cmax"].append(cmax_now)
            tr["arg"].append(float(np.argmax(C_new)))
            tr["c0"].append(float(C_new[0])); tr["cR"].append(float(C_new[-1]))
            tr["cmin"].append(cmin_now)

        # ---------------- 达标判定（首个 C_max ≤ 0.15；< 与 ≤ 的 inf 相等）----------------
        if stop_on_dry and t_dry is None and cmax_now <= c_dry:
            t_dry = t
            t_dry_prev = (t - dt_now, cmax_prev)
            t_end = t + margin_s
            n_end = math.ceil(t_end / dt_now - 1e-9)
            t_end = n_end * dt_now
            if verbose:
                print("[dry] 首达 C_max ≤ %g：t = %.6f s (%.4f h, %.4f d)；C_max=%.9f；"
                      "→ t_end = %.4f s" % (c_dry, t, t / 3600.0, t / 86400.0, cmax_now, t_end),
                      flush=True)

        # ---------------- 快照 ----------------
        while (snap_next is not None and t >= snap_next - 1e-9
               and snap_next <= snap_stop + 1e-9):
            snap_t.append(t); snap_T.append(T[sn].copy()); snap_C.append(C[sn].copy())
            snap_next += snap_every_s

        # ---------------- 全场剖面取样 ----------------
        while prof_left and t >= prof_left[0] - 1e-9:
            profiles[prof_left.pop(0)] = (T.copy(), C.copy())

        # ---------------- 检查点 ----------------
        for tc in checkpoints:
            if tc not in ckpt and t >= tc - 1e-9:
                ckpt[tc] = dict(t=t, T=T.copy(), C=C.copy(), cum_flux=cum_flux,
                                mass=mass, n_steps=step, floor_count=floor_mon.count,
                                dr_cm=dr_cm)

        if verbose and progress_every and step % progress_every == 0:
            n_recent = min(2000, picard_iters.n)
            print("[prog] t=%10.3f s (%.3f h) n=%8d C_max=%.9f C(0)=%.9f T(0)=%.6f "
                  "T(N)=%.6f it=%.2f floor=%d wall=%.1f s"
                  % (t, t / 3600.0, step, cmax_now, float(C_new[0]), float(T_new[0]),
                     float(T_new[N]), float(picard_iters.a[picard_iters.n - n_recent:
                                                          picard_iters.n].mean()),
                     floor_mon.count, time.time() - t_wall0), flush=True)
        cmax_prev = cmax_now

    iters = picard_iters.array()
    uniq, cnt = (np.unique(iters, return_counts=True) if iters.size else
                 (np.array([]), np.array([])))
    out = dict(
        status=status, params=P,
        r=r, dr=dr, N=N, W=W, dr_cm=dr_cm, geo_note="精确圆柱测度 W_0=Δr²/8, W_i=r_iΔr, "
                                                    "W_N=RΔr/2−Δr²/8",
        dt1=dt1, dt2=dt2, t_switch=t_switch, dt_final=dt_now, n_switch=n_switch,
        n_steps=n_steps_total, t_final=t, t_dry=t_dry, t_dry_prev=t_dry_prev,
        t_end=t_end, c_dry=c_dry,
        snap_t=np.array(snap_t), snap_T=np.array(snap_T), snap_C=np.array(snap_C),
        T_final=T.copy(), C_final=C.copy(),
        picard=dict(max_iters=int(iters.max()) if iters.size else 0,
                    mean_iters=float(iters.mean()) if iters.size else 0.0,
                    median_iters=float(np.median(iters)) if iters.size else 0.0,
                    n_not_converged=int(n_not_conv),
                    hist_iters=[int(v) for v in uniq],
                    hist_counts=[int(v) for v in cnt],
                    iters=iters),
        residual=dict(max_T=res_max_T, max_C=res_max_C,
                      note="‖A_φ φ − b_φ‖_∞（Q2-23 第 3 项），每内部步在最终迭代值上算一次"),
        picard_step_dmax=dict(T=dT_seen, C=dC_seen,
                              note="Picard 相邻迭代差的最大值（全时段）"),
        floor=floor_mon.to_dict(),
        violations=viol,
        audit=dict(n=aud["n"], max_abs=aud["max_abs"],
                   max_rel_perstep=aud["max_rel_perstep"],
                   max_rel_global=(aud["max_abs"] / max(aud["peak_dE"], aud["peak_flux"], 1e-300)),
                   mean_abs=(aud["sum_abs"] / aud["n"]) if aud["n"] else 0.0,
                   peak_dE=aud["peak_dE"], peak_flux=aud["peak_flux"],
                   samples=aud["samples"],
                   note=("式 (Q2-25)：ΣW_i[ω^{n+1}T^{n+1}−ω^nT^n] − ΣW_i T_i^n(ω^{n+1}−ω^n) "
                         "= Δt·边界通量 + O(Δt²)；对本离散格式是**代数恒等式**。"
                         "**判据必须用 max_rel_global**（= max|残差| / 全程 max(|ΔE|,|边界通量|)）；"
                         "max_rel_perstep 在后期 dE≈1e-9 的步上会被舍入噪声放大，单独看会误判")),
        mass=dict(m0=m0_ref, m1=mass,
                  ratio=mass / m0_ref if m0_ref else float("nan"), cum_flux=cum_flux,
                  rel_err=(abs((m0_ref - mass) - cum_flux) / abs(m0_ref - mass)
                           if m0_ref != mass else 0.0),
                  note="P3：Δ(ΣW_i C_i) 应等于表面累积流出（跨界面通量逐对相消，严格成立）"),
        cmax=dict(t=tr["t"].array() if cmax_track else None,
                  cmax=tr["cmax"].array() if cmax_track else None,
                  argmax=tr["arg"].array() if cmax_track else None,
                  c0=tr["c0"].array() if cmax_track else None,
                  cR=tr["cR"].array() if cmax_track else None,
                  cmin=tr["cmin"].array() if cmax_track else None),
        checkpoints=ckpt, profiles=profiles,
        elapsed_s=time.time() - t_wall0,
    )
    return out


# ============================================================================
# 5) 烟测（快速；不产生交付产物）
# ============================================================================
def smoke_test(dr_cm=0.01, dt=1.0 / 8.0, t_end=600.0, verbose=True):
    """短窗烟测：物理自检（升温/不超驱动/水分下降/剖面单调/无 floor 激活）。"""
    o = solve_q2(dr_cm=dr_cm, dt1=dt, dt2=dt, t_switch=1e18,
                 stop_on_dry=False, t_end_override=t_end,
                 snap_every_s=1.0, snap_until_s=60.0, cmax_track=True, verbose=False)
    T, C = o["T_final"], o["C_final"]
    ta, ca = driver_at(t_end)
    checks = {
        "T_center_rise": bool(T[0] > T_INIT_C),
        "T_surface_ge_center": bool(T[-1] >= T[0] - 1e-12),
        "T_not_exceed_driver": bool(T.max() <= ta + 1e-9),
        # 物理：表面最干、中心最湿 → C(r) 沿 r **非增**（∂C/∂r ≤ 0）
        "C_radial_nonincreasing": bool(np.all(np.diff(C) <= 1e-12)),
        "C_below_init": bool(C.max() <= C_INIT + 1e-12),
        "C_above_ca": bool(C.min() >= ca - 1e-12),
        "no_floor_activation": bool(o["floor"]["count"] == 0),
        "picard_all_converged": bool(o["picard"]["n_not_converged"] == 0),
        "no_extremum_violation": bool(len(o["violations"]) == 0),
    }
    if verbose:
        print("耗时 %.1f s；步数 %d；Picard 均值 %.3f 最大 %d 未收敛 %d"
              % (o["elapsed_s"], o["n_steps"], o["picard"]["mean_iters"],
                 o["picard"]["max_iters"], o["picard"]["n_not_converged"]))
        print("T: 中心 %.6f ℃  表面 %.6f ℃  驱动 %.6f ℃" % (T[0], T[-1], ta))
        print("C: 中心 %.9f  表面 %.9f  C_max %.9f  C_min %.9f" % (C[0], C[-1], C.max(), C.min()))
        print("D_floor 激活 %d 次；能量审计 max|残差| %.3e（全局相对 %.3e）；"
              "水分守恒相对误差 %.3e"
              % (o["floor"]["count"], o["audit"]["max_abs"],
                 o["audit"]["max_rel_global"], o["mass"]["rel_err"]))
        for k, v in checks.items():
            print("  %-24s %s" % (k, "PASS" if v else "**FAIL**"))
    return o, checks


if __name__ == "__main__":
    import platform
    print("=" * 78)
    print("[solver_q2] smoke test  dt=1/8 s, t_end=600 s, dr=0.01 cm")
    print("python %s / numpy %s" % (platform.python_version(), np.__version__))
    print("=" * 78)
    o, ck = smoke_test()
    print("=" * 78)
    print("全部自检通过" if all(ck.values()) else "**存在失败项**")
