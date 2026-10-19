# -*- coding: utf-8 -*-
"""
q3_selftest.py —— `src/q3_from_q2.py` 的**合成数据自测**（不依赖 Q2 真实结果）

为什么要有它
------------
Q3 搭 Q2 的车。Q2 求解还在跑，若等结果出来再写 Q3 脚本就多串行一段。
本脚本用**已知答案的合成 Q2 产物**把 Q3 的逻辑先验完：Q2 结果一落盘，
`q3_from_q2.py` 就能立刻出产物。

合成算例（全部落 `results/_q3_selftest_out/`，**绝不写 `results/` 根**）
----
  算例 1  解析已知答案：C_max(t) = 0.10 + 2.45·e^(−t/τ) → t_dry 可解析算出，
          检验**插值定位**能否复现该已知值（报实测误差）
  算例 2  端到端：合成 Q2 产物目录（轨迹 npz + results2.json + 全时段水分场 npz）
          → 跑完整 `run()` → 逐位核对 result3.xlsx 的形状/表头/sheet 名/数值，
          并核对论文表 5 的取数与末行「烘干结束时间」
  算例 3  边界①：C_max 轨迹**恰好**有一点等于 0.15 → t_dry 必须正是该点时刻
  算例 4  边界②：14 d 内未达标 → 必须报「未达标」，**不得静默返回末行**
  算例 5  边界③：输入文件缺失 → 必须明确报错，且列出「我期待什么 / 实际看到什么」

用法： python src/q3_selftest.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import numpy as np

for _s in (sys.stdout, sys.stderr):                 # Windows 重定向到文件/管道时默认 GBK 会崩
    if getattr(_s, "reconfigure", None):
        _s.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import q3_from_q2 as Q3                                            # noqa: E402

OUT_ROOT = ROOT / "results" / "_q3_selftest_out"
RESULTS = []

# 合成算例的口径（与 Q2 建模文档一致的内部步结构）
DT1, T_SWITCH, DT2 = 1.0 / 64, 14400.0, 1.0 / 8
TAU_1 = 58327.0                      # 算例 1：t_dry ≈ 63 h 量级（对齐 §7.3 的尺度估计）
TAU_2 = 18510.0                      # 算例 2：t_dry ≈ 20 h 量级（小规模，仍能出 6h 行）


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok), detail))
    print("  [%s] %s%s" % ("PASS" if ok else "FAIL", name,
                           ("  —— " + detail) if detail else ""), flush=True)
    return bool(ok)


def internal_grid(t_end, dt1=DT1, t_switch=T_SWITCH, dt2=DT2):
    """复刻 Q2 的两段式内部步时间轴（I 段细、II 段粗），步长为精确有理数。

    落盘的合成算例用较粗的 I 段（见各调用处的注释）：这些算例验的是**抽样/表头/错误路径**，
    不是定位精度（定位精度由算例 1 在内存里用真 Δt_I=1/64 验，不落盘）。
    真 Δt_I=1/64 的轨迹有 92 万点、单文件 36 MB —— 自测没必要为此占盘。
    """
    n1 = int(round(min(t_switch, t_end) / dt1))
    t1 = np.arange(1, n1 + 1) * dt1
    rest = t_end - n1 * dt1
    n2 = int(np.ceil(rest / dt2 - 1e-9)) if rest > 0 else 0
    return np.concatenate([t1, n1 * dt1 + np.arange(1, n2 + 1) * dt2])


def cmax_analytic(t, tau, floor=0.10, amp=2.45):
    """C_max(t) = floor + amp·e^(−t/τ)。首达 0.15 的解析解 t* = τ·ln(amp/(0.15−floor))。"""
    return floor + amp * np.exp(-np.asarray(t, float) / tau)


def t_dry_analytic(tau, level=Q3.C_DRY, floor=0.10, amp=2.45):
    return tau * np.log(amp / (level - floor))


def build_field(t_s, r_cm, tau, tau_s):
    """合成水分场：径向单调递减、C(0,t) = C_max(t)、C(R,t) = C_surf(t)。"""
    cmax = cmax_analytic(t_s, tau)
    csf = cmax * (0.6 + 0.4 * np.exp(-np.asarray(t_s, float) / tau_s))
    u = 1.0 - (r_cm[None, :] / r_cm[-1]) ** 2
    return cmax, csf[:, None] + (cmax - csf)[:, None] * u


def write_synth_dir(d: Path, t_s, cmax, field_t, r_cm, C, t_dry_json_s, status="ok",
                    t_end_s=None):
    """造一个**结构上等价于 Q2 产物**的合成目录。"""
    d.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(d / Q3.TRAJ_CANDIDATES[0], t_s=t_s, C_max=cmax,
                        argmax_i=np.zeros(t_s.size), C_center=cmax,
                        C_surface=cmax, C_min=cmax)
    np.savez_compressed(d / "q2_field_60s.npz", t_s=field_t, r_cm=r_cm, C_kgkg=C)
    (d / Q3.SUMMARY_CANDIDATES[0]).write_text(json_dumps({
        "meta": {"config": {"status": status}},
        "t_dry": {"value_s": float(t_dry_json_s),
                  "value_h": float(t_dry_json_s) / 3600.0,
                  "bracket": {"width_s": None},
                  "t_end_s": float(t_end_s if t_end_s is not None
                                   else t_dry_json_s + Q3.MARGIN_S),
                  "status": status},
    }), encoding="utf-8")
    return d


def json_dumps(o):
    import json
    return json.dumps(o, ensure_ascii=False, indent=2)


# --------------------------------------------------------------------------- #
# 算例 1：解析已知答案 → 检验插值定位
# --------------------------------------------------------------------------- #
def case1():
    print("\n=== 算例 1：解析已知答案（C_max(t) = 0.10 + 2.45·e^(−t/τ)，τ = %g s）==="
          % TAU_1)
    t_exact = t_dry_analytic(TAU_1)
    check("解析解量级合理（t_dry 落在 30–130 h，Q2 建模文档 §14.2 第 10 条的冒烟窗口）",
          30.0 <= t_exact / 3600.0 <= 130.0,
          "t_dry(解析) = %.6f s = %.6f h" % (t_exact, t_exact / 3600.0))

    t_s = internal_grid(t_exact + Q3.MARGIN_S)
    cm = cmax_analytic(t_s, TAU_1)
    loc = Q3.locate_t_dry(t_s, cm, Q3.C_DRY, Q3.T_CAP_S)

    check("定位器找到下穿点", loc["crossed"], "n* = %d" % loc.get("n_star"))
    err_s = abs(loc["t_linear_s"] - t_exact)
    err_rel = err_s / t_exact
    print("      实测：线性插值 t_dry = %.9f s；解析 t_dry = %.9f s；误差 = %.3e s "
          "(相对 %.2e)" % (loc["t_linear_s"], t_exact, err_s, err_rel))
    print("      括住区间 [%.6f, %.6f] s（宽 %.6f s）；C_max 两端 %.9f / %.9f"
          % (loc["t_before_s"], loc["t_at_cross_s"], loc["bracket_width_s"],
             loc["c_max_before"], loc["c_max_at_cross"]))
    print("      Hermite 值 %.9f s（与线性差 %.3e s）；跨越点 dC_max/dt = %.6e kg/(kg·s)"
          % (loc["t_hermite_s"], loc["localization_error_s"], loc["slope_at_cross_s"]))
    check("插值定位复现解析值（|Δ| ≤ 1e-3 s）", err_s <= 1e-3,
          "误差 %.3e s" % err_s)

    # 精度口径必须由**实测斜率**给出，而不是照抄 §11.3 的名义值。
    # 解析：C_max'(t) = −(2.45/τ)·e^(−t/τ)，在跨越点 e^(−t/τ) = 0.05/2.45 ⇒ |C'| = 0.05/τ。
    unc = Q3.sensitivity_from_slope(loc["slope_at_cross_s"])
    slope_exact = 0.05 / TAU_1
    # `slope_at_cross_*` 是**离散差商**（run_q2 用相邻步差商），精度 O(h) ——
    # 不是解析导数，故用相对容差比对，并把实测相对偏差打出来。
    rtol_a = abs(abs(loc["slope_at_cross_s"]) - slope_exact) / slope_exact
    check("跨越点斜率与解析值 0.05/τ 相符（离散差商，相对容差 1e-5）", rtol_a <= 1e-5,
          "实测 %.9e vs 解析 %.9e kg/(kg·s)，相对偏差 %.3e"
          % (abs(loc["slope_at_cross_s"]), slope_exact, rtol_a))
    check("敏感度 = 1/|dC_max/dt|，不确定度 = 1e-5 × 敏感度（内部恒等式，逐位）",
          abs(unc["factor_s_per_kgkg"] - 1.0 / abs(loc["slope_at_cross_s"])) <= 1e-9
          and abs(unc["t_dry_uncertainty_s"]
                  - Q3.ASSUMED_FIELD_ERROR_KGKG
                  / abs(loc["slope_at_cross_s"])) <= 1e-9,
          "敏感度 %.6e s per (kg/kg) → 不确定度 ±%.3f s = ±%.5f h"
          % (unc["factor_s_per_kgkg"], unc["t_dry_uncertainty_s"],
             unc["t_dry_uncertainty_h"]))
    print("      与 §11.3 名义口径的对照：名义 %.6e，实测 %.6e（实测为其 %.4f 倍）→ "
          "名义 ±%.1f s，实测 ±%.3f s"
          % (Q3.DOC_FIELD_TO_TIME_S_PER_KGKG, unc["factor_s_per_kgkg"],
             unc["factor_s_per_kgkg"] / Q3.DOC_FIELD_TO_TIME_S_PER_KGKG,
             Q3.DOC_T_DRY_UNCERTAINTY_S, unc["t_dry_uncertainty_s"]))
    check("定位误差远小于**实测**物理不确定度", err_s < unc["t_dry_uncertainty_s"],
          "%.3e s << %.3f s" % (err_s, unc["t_dry_uncertainty_s"]))
    check("斜率不可用（None/0）时退回 §11.3 名义口径而非崩溃",
          Q3.sensitivity_from_slope(None)["measured"] is False
          and Q3.sensitivity_from_slope(0.0)["t_dry_uncertainty_s"]
          == Q3.DOC_T_DRY_UNCERTAINTY_S,
          "None → ±%.1f s（名义）" % Q3.sensitivity_from_slope(None)["t_dry_uncertainty_s"])
    check("括住区间 = 1 个内部步（II 段 Δt = 1/8 s）",
          abs(loc["bracket_width_s"] - DT2) < 1e-12,
          "宽 %.6f s" % loc["bracket_width_s"])
    mono = Q3.monotonicity_check(t_s, cm, loc["n_star"])
    check("跨越前 C_max 非增（§11.2 单调性检查）", mono["n_rise_steps_over_tol"] == 0,
          mono["verdict"])
    check("未下穿时不报 crossed 之外的结论（轨迹若截断需另行提示）",
          loc["crossed"] and t_s[-1] > t_exact + Q3.MARGIN_S - 1e-6,
          "轨迹末端 %.1f s" % t_s[-1])
    return {"t_exact_s": t_exact, "t_lin_s": loc["t_linear_s"], "err_s": err_s}


# --------------------------------------------------------------------------- #
# 算例 2：端到端（合成 Q2 产物 → result3.xlsx + 表 5）
# --------------------------------------------------------------------------- #
def case2():
    print("\n=== 算例 2：端到端（合成 Q2 产物目录 → result3.xlsx + 论文表 5）===")
    t_exact = t_dry_analytic(TAU_2)
    d = OUT_ROOT / "case2" / "q2_like"
    odir = OUT_ROOT / "case2" / "out"
    if d.exists():
        shutil.rmtree(d)
    if odir.exists():
        shutil.rmtree(odir)

    # --- 合成 C_max 轨迹（逐内部步，Q2 结构） ---
    # I 段用 Δt=1/16 s（真 Q2 为 1/64 s）：本算例验抽样/表头/表 5，不是定位精度；
    # t_dry ≈ 72000 s 远在 II 段，跨越括住区间仍是 1/8 s —— 与真 Q2 同构。
    t_traj = internal_grid(t_exact + Q3.MARGIN_S, dt1=1.0 / 16)
    cm_traj = cmax_analytic(t_traj, TAU_2)
    # --- 合成全时段水分场（60 s 网格 × 21 列，覆盖 t_dry + 3 h 余量） ---
    t_field = np.arange(60.0, np.ceil((t_exact + Q3.MARGIN_S) / 60.0) * 60.0 + 1e-9, 60.0)
    _, C_field = build_field(t_field, Q3.R_OUT_CM, TAU_2, TAU_2 / 5.0)

    check("合成场为 60 s 网格（首点 60 s，步长 60 s）",
          t_field[0] == 60.0 and np.allclose(np.diff(t_field), 60.0),
          "%d 行 × %d 列，t = %.0f…%.0f s" % (C_field.shape[0], C_field.shape[1],
                                              t_field[0], t_field[-1]))
    check("合成场覆盖 t_dry（t_last > t_dry）", t_field[-1] > t_exact,
          "场末端 %.0f s ≥ t_dry %.3f s" % (t_field[-1], t_exact))
    # results2.json 里放**解析值**（不是本脚本复现值）→ 顺带检验交叉核对容差容得下离散化差
    write_synth_dir(d, t_traj, cm_traj, t_field, Q3.R_OUT_CM, C_field, t_exact)

    summary = Q3.run(d, odir, verbose=True)

    # ---- (a) t_dry 交叉核对 ----
    td = summary["t_dry"]
    print("      t_dry = %.6f s = %.6f h；Q2 自报 %.6f s；差 %.3e s"
          % (td["value_s"], td["value_h"], t_exact, td["source"]["crosscheck_abs_diff_s"]))
    check("与 Q2 自报值交叉核对通过（|Δ| ≤ 1e-3 s）",
          td["source"]["crosscheck_abs_diff_s"] <= 1e-3,
          "Δ = %.3e s" % td["source"]["crosscheck_abs_diff_s"])
    check("t_dry 复现解析值（|Δ| ≤ 1e-3 s）",
          abs(td["value_s"] - t_exact) <= 1e-3,
          "复现 %.6f vs 解析 %.6f，Δ = %.3e s"
          % (td["value_s"], t_exact, abs(td["value_s"] - t_exact)))
    check("h 单位换算正确（value_h == value_s/3600）",
          abs(td["value_h"] - td["value_s"] / 3600.0) < 1e-12,
          "%.6f h" % td["value_h"])

    # ---- (b) result3.xlsx 逐位核对 ----
    import openpyxl
    wb = openpyxl.load_workbook(odir / "result3.xlsx")
    ws = wb.active
    check("result3.xlsx sheet 名逐位等于附件 3 模板",
          wb.sheetnames == [Q3.SHEET_NAME_RESULT3],
          "sheets = %s（模板 = ['%s']）" % (wb.sheetnames, Q3.SHEET_NAME_RESULT3))
    hdr = [ws.cell(1, c).value for c in range(1, 23)]
    exp_hdr = [Q3.HEADER_A1] + [float(round(rc, 1)) for rc in Q3.R_OUT_CM]
    check("A1 表头 + 21 个半径表头逐位正确", hdr == exp_hdr,
          "A1 = %r；B1..U1 = %s" % (hdr[0], hdr[1:4]))
    check("列数 = 22（1 时间 + 21 半径）", ws.max_column == 22, "max_column = %d" % ws.max_column)

    t_last_out = float(np.floor(t_exact / 60.0 + 1e-9) * 60.0)
    times_exp = np.arange(60.0, t_last_out + 1e-9, 60.0)
    check("行数 = floor(t_dry/60)（含 60 s 首行、不含 t=0）",
          ws.max_row == times_exp.size + 1,
          "max_row = %d（表头 1 + 数据 %d）" % (ws.max_row, times_exp.size))
    t_col = np.array([ws.cell(r, 1).value for r in range(2, ws.max_row + 1)], float)
    check("A 列时间轴 = 60, 120, …, %d（逐位）" % int(t_last_out),
          np.array_equal(t_col, times_exp),
          "首 %.0f / 末 %.0f / 步长 %.0f" % (t_col[0], t_col[-1],
                                            np.unique(np.diff(t_col))[0]))

    idx, align_err = Q3.sample_on_grid(t_field, C_field, times_exp)
    got = np.array([[ws.cell(r, c).value for c in range(2, 23)]
                    for r in range(2, ws.max_row + 1)], float)
    exp = np.round(C_field[idx], 4)
    n_bad = int(np.count_nonzero(got != exp))
    maxdiff = float(np.abs(got - exp).max())
    check("result3 全部数值**逐位**等于 round(合成场, 4)", n_bad == 0,
          "%d×%d = %d 个格点，不等 %d 个，max|Δ| = %r" % (got.shape[0], got.shape[1],
                                                          got.size, n_bad, maxdiff))
    check("抽样对齐偏差为 0（60 s 是 60 s 的整数倍）", align_err == 0.0,
          "max|Δt| = %r s" % align_err)
    check("首行数值 = 合成场 t=60 s 行的四位舍入",
          np.allclose(got[0], np.round(C_field[0], 4), atol=0, rtol=0),
          "r=0/1.0/2.0 cm 处 %.4f / %.4f / %.4f" % (got[0][0], got[0][10], got[0][20]))

    # ---- (c) 论文表 5 ----
    txt = (odir / "q3_table5_moisture.csv").read_text(encoding="utf-8-sig").splitlines()
    rows = [r.split(",") for r in txt[1:]]
    t5 = summary["table5"]
    n_full_exp = int(np.floor(t_exact / 3600.0 / 6.0 + 1e-12))
    check("表 5 行数 = floor(t_dry_h/6) + 1（末行「烘干结束时间」）",
          len(rows) == n_full_exp + 1,
          "%d 行（6h 行 %d + 结束行 1）；t_dry_h = %.6f" % (len(rows), n_full_exp,
                                                           t_exact / 3600.0))
    check("表 5 列数 = 6（时间 + 5 个半径）", all(len(r) == 6 for r in rows),
          "各行列数 = %s" % sorted({len(r) for r in rows}))
    check("表 5 行标签 = 6/12/…/烘干结束时间",
          [r[0] for r in rows] == ["%g" % (6.0 * (k + 1)) for k in range(n_full_exp)]
          + ["烘干结束时间"],
          "%s" % [r[0] for r in rows])
    check("表 5 行标签与 JSON 的 labels 一致",
          [r[0] for r in rows] == t5["labels"],
          "%s" % t5["labels"])

    # 6h 行的数值：应在 60 s 网格上精确命中
    idx6, _ = Q3.sample_on_grid(t_field, C_field, np.array(t5["times_h"]) * 3600.0)
    i5 = [int(np.argmin(np.abs(Q3.R_OUT_CM - r))) for r in Q3.R_TAB5_CM]
    exp6 = np.round(C_field[idx6][:, i5], 4)
    got6 = np.array([[float(v) for v in r[1:]] for r in rows[:n_full_exp]])
    check("表 5 的 6h 行数值逐位正确",
          n_full_exp == 0 or np.array_equal(got6, exp6),
          "r=0 cm 列 %s" % (got6[:, 0].tolist() if n_full_exp else "（无 6h 行）"))
    # 末行：t_dry 处线性插值
    row_end, brk = Q3.interp_linear_columns(t_field, C_field, t_exact)
    exp_end = np.round(row_end[i5], 4)
    got_end = np.array([float(v) for v in rows[-1][1:]])
    check("表 5「烘干结束时间」行数值逐位正确（t_dry 处线性插值）",
          np.array_equal(got_end, exp_end),
          "得到 %s；合成场插值 %s" % (got_end.tolist(), exp_end.tolist()))
    check("末行 r=0 cm 处 ≈ 0.15（判据成立：中心恰好压线）",
          abs(got_end[0] - Q3.C_DRY) <= 5e-4,
          "C(0, t_dry) = %.6f（判据 %.2f）；括住区间 [%.0f, %.0f] s"
          % (got_end[0], Q3.C_DRY, brk[0], brk[1]))
    exc_s = summary["checks"]["row_at_t_dry_center_excess_time_s"]
    unc5 = summary["t_dry"]["uncertainty"]
    # 2026-09-11 口径改造后，summary 的 uncertainty 是分层结构：
    # 主报网格项 = grid_contribution_s；灵敏度 = field_error_sensitivity_s_per_kgkg。
    # 旧键 t_dry_uncertainty_s / factor_s_per_kgkg 只留在 deprecated 块里，不再作为主报项。
    grid_unc_s = unc5.get("grid_contribution_s")
    sens_s = unc5.get("field_error_sensitivity_s_per_kgkg")
    if grid_unc_s is None:
        print("  [SKIP] 末行中心列的超量换算成时间当量后落在主报网格项内"
              "（合成算例未提供同源的 q2_error_budget.json，主报网格项待回填）")
    else:
        check("末行中心列的超量换算成时间当量后落在主报网格项内（弦曲率效应，非判据违反）",
              abs(exc_s) <= float(grid_unc_s),
              "超量 %+.3e kg/kg = %+.4f s 时间当量（主报网格项 ±%.3f s）"
              % (summary["checks"]["row_at_t_dry_center_excess_kgkg"], exc_s,
                 float(grid_unc_s)))
    rtol2 = abs(sens_s - TAU_2 / 0.05) / (TAU_2 / 0.05) if sens_s is not None else float("inf")
    check("算例 2 的敏感度也由实测斜率给出（≈ τ/0.05，离散差商相对容差 1e-5）",
          rtol2 <= 1e-5,
          "实测 %s vs 解析 τ/0.05 = %.6e s per (kg/kg)，相对偏差 %.3e"
          % ("待回填" if sens_s is None else "%.6e" % sens_s, TAU_2 / 0.05, rtol2))
    check("末行沿半径单调不增（r=0 最大、r=2 最小）",
          bool(np.all(np.diff(got_end) <= 0)),
          "%s" % got_end.tolist())
    check("表 5 的列半径 = 0/0.5/1/1.5/2 cm",
          t5["r_cm"] == [0.0, 0.5, 1.0, 1.5, 2.0], "%s" % t5["r_cm"])

    # ---- (d) 溯源与安全 ----
    check("q3_t_dry.json / run_manifest_q3.json 已落盘",
          (odir / "q3_t_dry.json").exists() and (odir / "run_manifest_q3.json").exists(),
          "q3_t_dry.json %d B"
          % (odir / "q3_t_dry.json").stat().st_size)
    check("产物名与 Q1/Q2 产物无交集（未覆盖任何 Q1/Q2 产物）",
          not (set(Q3.Q3_OUTPUT_NAMES) & Q3.Q12_OWNED),
          "Q3 产物 = %s" % list(Q3.Q3_OUTPUT_NAMES))
    check("自测未在 results/ 根写 result3.xlsx", not (ROOT / "results" / "result3.xlsx")
          .exists() or PREEXIST_RESULT3,
          "results/result3.xlsx 存在=%s（自测前存在=%s）"
          % ((ROOT / "results" / "result3.xlsx").exists(), PREEXIST_RESULT3))
    return summary


# --------------------------------------------------------------------------- #
# 算例 3：边界① C_max 恰好等于 0.15
# --------------------------------------------------------------------------- #
def case3():
    print("\n=== 算例 3（边界①）：C_max 轨迹**恰好**有一点等于 0.15 ===")
    T = 100000.0                                   # 恰在内部步网格上
    t_s = np.arange(1, int(T / DT2) + 4000 + 1) * DT2
    cm = np.where(t_s <= T, Q3.C_DRY + 2.40 * (1.0 - t_s / T), Q3.C_DRY)
    n_hit = int(np.count_nonzero(cm == Q3.C_DRY))
    check("合成轨迹确实存在严格等于 0.15 的点", n_hit > 0,
          "%d 个点（首个 t = %.6f s）"
          % (n_hit, float(t_s[np.flatnonzero(cm == Q3.C_DRY)[0]])))
    loc = Q3.locate_t_dry(t_s, cm, Q3.C_DRY, Q3.T_CAP_S)
    check("定位结果 = 恰好压线的那一刻（|Δ| ≤ 1e-9 s）",
          abs(loc["t_linear_s"] - T) <= 1e-9,
          "t_linear = %.9f s（期望 %.9f s，Δ = %.3e s）"
          % (loc["t_linear_s"], T, abs(loc["t_linear_s"] - T)))
    check("括住区间 = [T−Δt, T]（前一时刻仍在 0.15 之上）",
          abs(loc["t_before_s"] - (T - DT2)) < 1e-9
          and abs(loc["t_at_cross_s"] - T) < 1e-9,
          "[%.6f, %.6f]；C_max = %.9f → %.9f"
          % (loc["t_before_s"], loc["t_at_cross_s"],
             loc["c_max_before"], loc["c_max_at_cross"]))
    print("      注：Hermite 值 %.9f s 偏离 %.3e s —— 因压线点之后的平台段把 m1 差商拉平"
          % (loc["t_hermite_s"], loc["localization_error_s"]))
    check("`<` 与 `≤` 的 inf 相等：取首个 ≤ 0.15 的步即得压线时刻",
          abs(loc["t_linear_s"] - T) <= 1e-9, "两者同为 %.6f s" % T)
    return {"T_s": T, "t_lin_s": loc["t_linear_s"]}


# --------------------------------------------------------------------------- #
# 算例 4：边界② 14 d 内未达标
# --------------------------------------------------------------------------- #
def case4():
    print("\n=== 算例 4（边界②）：14 d 内未达标 ===")
    tau = 2.0e5
    t_s = np.arange(1, int(Q3.T_CAP_S / 60.0) + 1) * 60.0          # 60 s 网格到 14 d
    cm = 1.0 + 1.55 * np.exp(-t_s / tau)                            # 下确界 1.0 > 0.15
    loc = Q3.locate_t_dry(t_s, cm, Q3.C_DRY, Q3.T_CAP_S)
    check("定位器判定未下穿（crossed=False）", not loc["crossed"],
          "轨迹 %d 点，t 到 %.1f s（%.3f d）" % (loc["n_steps"], loc["t_last_s"],
                                                loc["t_last_s"] / 86400.0))
    check("正确识别「轨迹已到 t_cap」", bool(loc["reached_t_cap"]),
          "t_last = %.1f s ≥ t_cap = %.1f s" % (loc["t_last_s"], Q3.T_CAP_S))
    check("报出末点与全程最小 C_max（供论文说明缺口）",
          abs(loc["c_max_last"] - float(cm[-1])) < 1e-12
          and abs(loc["c_max_min"] - float(cm.min())) < 1e-12,
          "末点 %.6f，全程最小 %.6f（判据 %.2f）"
          % (loc["c_max_last"], loc["c_max_min"], Q3.C_DRY))

    d = OUT_ROOT / "case4" / "q2_like"
    if d.exists():
        shutil.rmtree(d)
    write_synth_dir(d, t_s, cm, t_s, Q3.R_OUT_CM,
                    np.tile(cm[:, None], (1, 21)), float(cm[-1]))
    try:
        Q3.run(d, OUT_ROOT / "case4" / "out", verbose=False)
        check("run() 对未达标**必须报错**而非静默返回末行", False,
              "run() 未报错就返回了 —— 违反 §11.4")
    except Q3.Q3InputError as e:
        msg = str(e)
        check("run() 报「未达标」并拒绝出声", "未达标" in msg,
              msg.splitlines()[0])
        check("报错信息含判据、轨迹范围与末点 C_max",
              all(k in msg for k in ("0.15", "t_cap", "末点 C_max")),
              " / ".join(l.strip() for l in msg.splitlines()[1:4]))
    out_dir = OUT_ROOT / "case4" / "out"
    check("未达标时**不得**产出 result3.xlsx",
          not (out_dir / "result3.xlsx").exists(),
          "out 目录存在=%s（%s）" % (out_dir.exists(),
                                     sorted(p.name for p in out_dir.iterdir())
                                     if out_dir.exists() else "空"))
    return {"t_last_s": loc["t_last_s"], "c_max_min": loc["c_max_min"]}


# --------------------------------------------------------------------------- #
# 算例 5：边界③ 输入文件缺失
# --------------------------------------------------------------------------- #
def case5():
    print("\n=== 算例 5（边界③）：输入文件缺失 ===")
    d = OUT_ROOT / "case5" / "q2_like_empty"
    if d.exists():
        shutil.rmtree(d)
    d.mkdir(parents=True, exist_ok=True)                  # 空目录：什么产物都没有
    try:
        Q3.run(d, OUT_ROOT / "case5" / "out", verbose=False)
        check("空目录必须报错（不得静默跳过）", False, "run() 未报错就返回了")
    except Q3.Q3InputError as e:
        msg = str(e)
        check("报错标题明确指向缺失的必需输入",
              "找不到 Q2 的 C_max 逐内部步轨迹" in msg, msg.splitlines()[0])
        check("列出「我期待什么」", "我期待：" in msg,
              [l.strip() for l in msg.splitlines() if "我期待" in l][0])
        check("列出「我实际看到什么」（含缺失路径与检索目录）",
              "我实际看到：" in msg and "缺失" in msg and "检索目录" in msg,
              " / ".join(l.strip() for l in msg.splitlines() if "缺失" in l
                         or "检索目录" in l))
        check("报错含 Q2 候选文件名与所需键名",
              Q3.TRAJ_CANDIDATES[0] in msg and "C_max" in msg,
              "候选 = %s" % list(Q3.TRAJ_CANDIDATES))

    # 半个输入：有轨迹（且**确实越过判据**），缺水分场 → 必须报「水分场缺失」
    d2 = OUT_ROOT / "case5" / "q2_like_no_field"
    if d2.exists():
        shutil.rmtree(d2)
    t_exact2 = t_dry_analytic(TAU_2)
    # 本子例只验「轨迹在、水分场不在」的错误路径 → 单段 1 s 网格足够（不验精度）
    t_s = internal_grid(t_exact2 + Q3.MARGIN_S, dt1=1.0, t_switch=0.0, dt2=1.0)
    cm = cmax_analytic(t_s, TAU_2)
    d2.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(d2 / Q3.TRAJ_CANDIDATES[0], t_s=t_s, C_max=cm)
    (d2 / Q3.SUMMARY_CANDIDATES[0]).write_text(json_dumps({
        "meta": {"config": {"status": "ok"}},
        "t_dry": {"value_s": float(t_exact2), "t_end_s": float(t_exact2 + Q3.MARGIN_S),
                  "status": "ok"}}), encoding="utf-8")
    try:
        Q3.run(d2, OUT_ROOT / "case5" / "out2", verbose=False)
        check("只有轨迹、没有全时段水分场时必须报错", False, "run() 未报错就返回了")
    except Q3.Q3InputError as e:
        msg = str(e)
        check("报错指向「覆盖 t_dry 的全时段水分场」缺失",
              "全时段水分场" in msg, msg.splitlines()[0])
        check("报错列出各候选的核验结果（缺失/跨度不足）", "缺失：" in msg,
              " / ".join(l.strip() for l in msg.splitlines()
                         if l.strip().startswith("- 缺失")))

    # Q2 因 14 d 未达标而中止：**落产物之前**就退出 → 目录里只有日志，没有轨迹。
    # 此时必须把日志里的「未达标」痕迹指出来，不能只报「文件缺失」（§11.4）。
    d4 = OUT_ROOT / "case5" / "q2_like_aborted"
    if d4.exists():
        shutil.rmtree(d4)
    d4.mkdir(parents=True, exist_ok=True)
    (d4 / "_run_q2.log").write_text(
        "[run_q2] 生产运行：Δr=0.01 cm, Δt_I=1/64 s\n"
        "[FATAL] 到达硬上限 t_cap = 14 d 仍未满足 C_max ≤ 0.15 —— "
        "按 §11.4 必须报「未达标」，不得静默返回末行\n"
        "SystemExit: 1\n", encoding="utf-8")
    try:
        Q3.run(d4, OUT_ROOT / "case5" / "out4", verbose=False)
        check("Q2 未达标中止（目录只有日志）时必须报错", False, "run() 未报错就返回了")
    except Q3.Q3InputError as e:
        msg = str(e)
        check("报错指出「14 d 未达标」痕迹，而非只报文件缺失",
              "未达标" in msg and "痕迹" in msg,
              " / ".join(l.strip() for l in msg.splitlines() if "痕迹" in l))

    # 跨度不足的专项：场只到 10800 s（= Q2 现有 result2 快照窗）
    d3 = OUT_ROOT / "case5" / "q2_like_short_field"
    if d3.exists():
        shutil.rmtree(d3)
    t_field = np.arange(60.0, 10800.0 + 1e-9, 60.0)
    _, C_field = build_field(t_field, Q3.R_OUT_CM, TAU_2, TAU_2 / 5.0)
    # 本子例只验「水分场跨度不足」的错误路径 → 单段 1 s 网格足够（不验精度）
    t_traj = internal_grid(t_dry_analytic(TAU_2) + Q3.MARGIN_S,
                           dt1=1.0, t_switch=0.0, dt2=1.0)
    write_synth_dir(d3, t_traj, cmax_analytic(t_traj, TAU_2), t_field, Q3.R_OUT_CM,
                    C_field, t_dry_analytic(TAU_2))
    try:
        Q3.run(d3, OUT_ROOT / "case5" / "out3", verbose=False)
        check("水分场跨度不足 t_dry 时必须报错（不得截断出表）", False,
              "run() 未报错就返回了")
    except Q3.Q3InputError as e:
        msg = str(e)
        check("报错明确写出「跨度不足」与缺多少小时", "跨度不足" in msg,
              " / ".join(l.strip() for l in msg.splitlines()
                         if "跨度不足" in l)[:200])
    return {}


# --------------------------------------------------------------------------- #
def main():
    global PREEXIST_RESULT3
    PREEXIST_RESULT3 = (ROOT / "results" / "result3.xlsx").exists()
    print("=" * 78)
    print("[q3_selftest] 合成数据自测（**不依赖 Q2 真实结果**）")
    print("  输出根：%s" % OUT_ROOT.as_posix())
    print("  （不写 results/ 根目录）")
    print("=" * 78)
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    c1 = case1()
    case2()
    c3 = case3()
    c4 = case4()
    case5()

    n_fail = sum(1 for _, ok, _ in RESULTS if not ok)
    print("\n" + "=" * 78)
    print("[汇总] %d 项检查：通过 %d，失败 %d" % (len(RESULTS), len(RESULTS) - n_fail,
                                               n_fail))
    for name, ok, detail in RESULTS:
        if not ok:
            print("  FAIL: %s%s" % (name, ("  —— " + detail) if detail else ""))
    print("---- 关键数字 ----")
    print("  算例 1  解析 t_dry = %.6f s；插值定位 %.6f s；误差 %.3e s"
          % (c1["t_exact_s"], c1["t_lin_s"], c1["err_s"]))
    print("  算例 3  压线时刻 %.6f s；定位 %.6f s" % (c3["T_s"], c3["t_lin_s"]))
    print("  算例 4  末点 %.1f s（%.3f d）；全程最小 C_max %.6f（未达标）"
          % (c4["t_last_s"], c4["t_last_s"] / 86400.0, c4["c_max_min"]))
    print("=" * 78)
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
