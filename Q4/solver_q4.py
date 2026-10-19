from __future__ import annotations

import math
import os
import sys
import time
from pathlib import Path

import numpy as np
from scipy.linalg.lapack import dgtsv

_OUT = sys.stdout
if hasattr(_OUT, "reconfigure"):
    _OUT.reconfigure(encoding="utf-8", errors="replace")

_BASE = Path(__file__).resolve().parent.parent

RHO_A4, RHO_B4 = 760.0, 90.0
CP_A4, CP_B4 = 1850.0, 2150.0
K_A4, K_B4 = 0.12, 0.20
D0_AP4, BETA_AP4, E_D_AP4 = 4.2e-4, 0.30, 3850.0

A2_RHO, A2_CP, A2_K = 820.0, 2600.0, 0.36
A2_D0, A2_BETA = 7.0e-9, 0.89
LOG_A2_D0 = math.log(A2_D0)

H_DEFAULT = 25.0
HM_DEFAULT = 8.0e-7

R_CM = 2.0
R0 = R_CM / 100.0
R_MIN_CM = 1.198
R_MIN = R_MIN_CM / 100.0
LENGTH_CM = 25.0
T_INIT_C = 28.0
C_INIT = 2.55

T_FREEZE_S = 14400.0
TA_FREEZE_C = 50.165
CA_FREEZE = 0.04986
TA_MAX_C = 50.246
CA_MIN = 0.01963

R_FREEZE_T = 259200.0
R_FREEZE = R_MIN

C_DRY = 0.15
MARGIN_S = 10800.0
T_CAP_S = 14.0 * 86400.0

D_FLOOR = 1.0e-25
LOG_D_FLOOR = math.log(D_FLOOR)
C_UNDERFLOW = 0.30 / 745.1332
LOG_D0_AP4 = math.log(D0_AP4)

D_AP4_2P55_323K = 2.5002509122378996e-9

def find_data_dir() -> Path:
    if os.environ.get("MM_DATA_DIR"):
        return Path(os.environ["MM_DATA_DIR"])
    for _c in (_BASE, _BASE / "A题", _BASE / "附件", _BASE.parent / "A题",
               Path.cwd() / "A题", Path.cwd()):
        if (_c / "附件" / "附件1.xlsx").is_file():
            return _c
    raise SystemExit(
        "未找到 A 题数据目录（须含 附件/附件1.xlsx）：请设置环境变量 MM_DATA_DIR 指向该目录，"
        "或把「A题」目录放到项目根/项目同级。"
    )


_DATA_DIR = str(find_data_dir())
ATT1 = str(Path(_DATA_DIR) / "附件" / "附件1.xlsx")
ATT2 = str(Path(_DATA_DIR) / "附件" / "附件2.xlsx")
ATT3_DIR = Path(_DATA_DIR) / "附件" / "附件3"
RESULT4_TEMPLATE = str(ATT3_DIR / "result4.xlsx")
PDF = str(Path(_DATA_DIR) / "A题.pdf")

_ATT1_CACHE = None
_ATT2_CACHE = None
_DRIVER_CACHE = {}


def att1():
    global _ATT1_CACHE
    if _ATT1_CACHE is None:
        import openpyxl
        wb = openpyxl.load_workbook(ATT1, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        t = np.array([float(r[0]) for r in rows])
        T = np.array([float(r[1]) for r in rows])
        C = np.array([float(r[2]) for r in rows])
        _ATT1_CACHE = (t, T, C)
    return _ATT1_CACHE


def att2():
    global _ATT2_CACHE
    if _ATT2_CACHE is None:
        import openpyxl
        wb = openpyxl.load_workbook(ATT2, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        t = np.array([float(r[0]) for r in rows])
        R_cm = np.array([float(r[1]) for r in rows])
        _ATT2_CACHE = (t, R_cm / 100.0)
    return _ATT2_CACHE


def driver_table():
    global _DRIVER_CACHE
    if not _DRIVER_CACHE:
        t, T, C = att1()
        dt_s = float(t[1] - t[0])
        nseg = int(round(t[-1] / dt_s))
        if abs(t[0]) > 1e-12 or abs(t[-1] - T_FREEZE_S) > 1e-9:
            raise ValueError("附件 1 时间轴不是 0…14400 s（步长 %g）" % dt_s)
        _DRIVER_CACHE = dict(dt_s=dt_s, nseg=nseg, T=T.copy(), C=C.copy(),
                             inv_dt=1.0 / dt_s, Ta_max=float(T.max()),
                             Ca_min=float(C.min()), Ca_max=float(C.max()))
    return _DRIVER_CACHE


def driver_at(t, interp="linear"):
    if t >= T_FREEZE_S:
        return TA_FREEZE_C, CA_FREEZE
    d = driver_table()
    x = t * d["inv_dt"]
    i = int(x)
    if i >= d["nseg"]:
        i = d["nseg"] - 1
    if interp == "step":
        return float(d["T"][i]), float(d["C"][i])
    f = x - i
    T = float(d["T"][i] + (d["T"][i + 1] - d["T"][i]) * f)
    C = float(d["C"][i] + (d["C"][i + 1] - d["C"][i]) * f)
    return T, C


def radius_at(t, mode="shrink"):
    arr = np.asarray(t, float)
    if mode == "const":
        out = np.full(arr.shape, R0)
    elif mode == "pchip":
        from scipy.interpolate import PchipInterpolator
        tt, RR = att2()
        out = PchipInterpolator(tt, RR, extrapolate=False)(np.clip(arr, 0.0, R_FREEZE_T))
        out = np.where(arr <= 0.0, R0, out)
        out = np.where(arr >= R_FREEZE_T, R_FREEZE, out)
    else:
        tt, RR = att2()
        out = np.interp(arr, tt, RR)
        out = np.where(arr <= 0.0, R0, out)
        out = np.where(arr >= R_FREEZE_T, R_FREEZE, out)
    return out if np.ndim(t) else float(out)


def radius_check(mode="shrink", n=20001):
    tt = np.linspace(0.0, R_FREEZE_T * 1.5, n)
    RR = radius_at(tt, mode)
    return dict(R_min=float(RR.min()), R_max=float(RR.max()),
                mono_nonincreasing=bool(np.all(np.diff(RR) <= 1e-15)),
                monotone_ok=bool(RR.min() >= R_MIN - 1e-12 and RR.max() <= R0 + 1e-12),
                n_samples=int(n))


def default_params():
    return dict(
        RHO_A=RHO_A4, RHO_B=RHO_B4, CP_A=CP_A4, CP_B=CP_B4, K_A=K_A4, K_B=K_B4,
        D0=D0_AP4, BETA=BETA_AP4, E_D=E_D_AP4, LOG_D0=LOG_D0_AP4,
        A2_RHO=A2_RHO, A2_CP=A2_CP, A2_K=A2_K, A2_D0=A2_D0, A2_BETA=A2_BETA,
        LOG_A2_D0=LOG_A2_D0,
        H=H_DEFAULT, HM=HM_DEFAULT,
        props="app4",
        D_CONST=None,
        RHO_SCALE=1.0, CP_SCALE=1.0,
        T_INIT=T_INIT_C, C_INIT=C_INIT,
    )


def rho_of_C(C, P):
    if P["props"] == "app2":
        return np.full_like(C, P["A2_RHO"])
    return P["RHO_A"] + P["RHO_B"] * C


def cp_of_C(C, P):
    if P["props"] == "app2":
        return np.full_like(C, P["A2_CP"])
    return P["CP_A"] + P["CP_B"] * C / (C + 1.0)


def omega_of_C(C, P):
    sc = P.get("RHO_SCALE", 1.0) * P.get("CP_SCALE", 1.0)
    base = rho_of_C(C, P) * cp_of_C(C, P)
    return base if sc == 1.0 else base * sc


def k_of_C(C, P):
    if P["props"] == "app2":
        return np.full_like(C, P["A2_K"])
    return P["K_A"] + P["K_B"] * C / (C + 1.0)


def logD_of_CT(C, T_C, P):
    if P.get("D_CONST") is not None:
        return np.full_like(np.asarray(C, float), math.log(P["D_CONST"]))
    if P["props"] == "app2":
        return P["LOG_A2_D0"] - P["A2_BETA"] / C
    T_K = T_C + 273.15
    return P["LOG_D0"] - P["BETA"] / C - P["E_D"] / T_K


class FloorMonitor:

    def __init__(self):
        self.count = 0
        self.events = []
        self.max_events = 20

    def note(self, mask, logD, C, T_C, t):
        n = int(np.count_nonzero(mask))
        if n == 0:
            return
        self.count += n
        Tc = np.asarray(T_C, float)
        for i in np.flatnonzero(mask):
            if len(self.events) >= self.max_events:
                break
            self.events.append(dict(t_s=float(t), i=int(i), C=float(C[i]),
                                    T_C=float(Tc[i]), logD=float(logD[i]),
                                    logD_floor=LOG_D_FLOOR))

    def to_dict(self):
        return dict(count=int(self.count), log_D_floor=LOG_D_FLOOR, D_floor=D_FLOOR,
                    c_underflow_threshold=C_UNDERFLOW,
                    note=("激活次数必须为 0（§5.3：极值原理给 C ≥ min(C_a)=%r，"
                          "是下溢点 %r 的 48.8 倍）。> 0 表示前提被推翻  停止并回报。"
                          % (CA_MIN, C_UNDERFLOW)),
                    first_events=self.events)


def eval_D(logD, monitor, C, T_C, t):
    if float(logD.min()) < LOG_D_FLOOR:
        monitor.note(logD < LOG_D_FLOOR, logD, C, T_C, t)
        logD = np.maximum(logD, LOG_D_FLOOR)
    return np.exp(logD)


def interface_gam(node, rule):
    if rule == "arith":
        return 0.5 * (node[:-1] + node[1:])
    if rule == "harm":
        a, b = node[:-1], node[1:]
        return 2.0 * a * b / (a + b)
    raise ValueError("未知界面系数规则 face_rule=%r" % (rule,))


def face_gam(kind, node, C, T_C, rule, P):
    if rule == "faceCT":
        Cf = 0.5 * (C[:-1] + C[1:])
        Tf = 0.5 * (T_C[:-1] + T_C[1:])
        if kind == "k":
            return k_of_C(Cf, P)
        return np.exp(np.maximum(logD_of_CT(Cf, Tf, P), LOG_D_FLOOR))
    return interface_gam(node, rule)


def build_grid(N):
    dxi = 1.0 / N
    xi = np.arange(N + 1) * dxi
    W = np.empty(N + 1)
    W[0] = dxi * dxi / 8.0
    W[1:N] = xi[1:N] * dxi
    W[N] = dxi / 2.0 - dxi * dxi / 8.0
    return xi, dxi, N, W


def make_geo(N, dxi, W):
    ii = np.arange(1, N)
    xip = ii * dxi + dxi / 2.0
    xim = ii * dxi - dxi / 2.0
    return dict(N=N, dxi=dxi, W=W,
                G1=xip / (dxi * W[1:N]),
                G2=xim / (dxi * W[1:N]),
                AXC=0.5 / W[0],
                SFC=(1.0 - dxi / 2.0) / (dxi * W[N]),
                SBND=1.0 / W[N])


def tdsolve(d, l, u, rhs):
    n = rhs.size
    x = dgtsv(l[1:], d, u[:-1], rhs.reshape(n, 1), overwrite_b=1)[3]
    return x[:, 0]


def resid_of(m, phi):
    d, l, u, rhs = m
    return banded_residual(d, l, u, phi, rhs)


def banded_residual(d, l, u, phi, rhs):
    res = d * phi
    res[:-1] += u[:-1] * phi[1:]
    res[1:] += l[1:] * phi[:-1]
    res -= rhs
    return float(np.abs(res).max())


def assemble(g, geo, gam_face, cap, om_extra_h, R_new, bc, rhs_old, driver_new):
    N = geo["N"]
    d = np.empty(N + 1); l = np.zeros(N + 1); u = np.zeros(N + 1)
    a1 = g * geo["G1"] * gam_face[1:N] / cap[1:N]
    a2 = g * geo["G2"] * gam_face[0:N - 1] / cap[1:N]
    d[1:N] = 1.0 + a1 + a2; u[1:N] = -a1; l[1:N] = -a2
    a0 = g * geo["AXC"] * gam_face[0] / cap[0]
    d[0] = 1.0 + a0; u[0] = -a0
    alpha = g * geo["SFC"] * gam_face[N - 1] / cap[N]
    beta = g * geo["SBND"] * R_new * om_extra_h / cap[N]
    d[N] = 1.0 + alpha + beta; l[N] = -alpha
    rhs = rhs_old.copy()
    rhs[N] += beta * driver_new
    if bc == "dirichlet":
        d[N] = 1.0; l[N] = 0.0; rhs[N] = driver_new
    rhs_keep = rhs.copy()
    phi = tdsolve(d, l, u, rhs)
    return phi, (d, l, u, rhs_keep)


class Grow:

    def __init__(self, cap=4096):
        self.a = np.empty(cap)
        self.n = 0

    def append(self, v):
        if self.n == self.a.size:
            self.a = np.resize(self.a, self.a.size * 2)
        self.a[self.n] = v
        self.n += 1

    def array(self):
        return self.a[:self.n].copy()


def scenario_driver(scen, samples, const):
    s = np.asarray(samples, float)
    if scen == "freeze_last":
        return s.copy()
    if scen == "const":
        return np.full_like(s, float(const))
    tail = s[-60:]
    if scen == "tail_mean":
        return np.concatenate([s[:-60], np.full(60, float(tail.mean()))])
    if scen == "row_min":
        return np.concatenate([s[:-60], np.full(60, float(tail.min()))])
    if scen == "row_max":
        return np.concatenate([s[:-60], np.full(60, float(tail.max()))])
    raise ValueError("未知情景 %r" % (scen,))


def solve_q4(N=200, dt1=5.0, dt2=20.0, t_switch=T_FREEZE_S,
             R_mode="shrink", geo_layer="nplus1", face_rule="arith",
             props="app4", bc="robin", T_mode="prescribed", block_order="TC",
             d_const=None, t_cap=T_CAP_S, margin_s=MARGIN_S, c_dry=C_DRY,
             stop_on_dry=True, t_end_override=None, t_start_state=None,
             picard_tol=1e-12, picard_max=100,
             snap_every_s=60.0, snap_until_s=None, snap_end=None, snap_nodes=None,
             profile_times=(), cmax_track=True, checkpoints=(), audit=True,
             audit_series=False,
             driver_interp="linear", ta_scenario="freeze_last",
             ca_scenario="freeze_last", const_ta=TA_FREEZE_C, const_ca=CA_FREEZE,
             verbose=True, progress_every=0, **overrides):
    Ca_scen, N, P, Ta_scen, W, _inv_dt, _nseg, dxi, geo, t_wall0, use_frozen_tbl, xi = setup_params(N, R_mode, T_mode, bc, ca_scenario, const_ca, const_ta, d_const, dt1, dt2, geo_layer, overrides, props, snap_every_s, ta_scenario)

    c_hi, c_lo, drv_at, t_hi, t_lo = prepare_drivers(Ca_scen, P, Ta_scen, _inv_dt, _nseg, driver_interp, use_frozen_tbl)

    t = 0.0
    step = 0
    cum_flux = 0.0
    if t_start_state is not None:
        t = float(t_start_state["t"])
        step = int(t_start_state.get("n_steps", 0))
        T = np.array(t_start_state["T"], float)
        C = np.array(t_start_state["C"], float)
        cum_flux = float(t_start_state.get("cum_flux", 0.0))
        if T.size != N + 1:
            raise ValueError("检查点网格与 N=%d 不符" % N)
    else:
        T = np.full(N + 1, float(P["T_INIT"]))
        C = np.full(N + 1, float(P["C_INIT"]))

    R_now = float(radius_at(t, R_mode))
    ta_old = float(drv_at(t)[0])
    floor_mon = FloorMonitor()
    viol = []

    if snap_until_s is None and snap_end is None:
        snap_stop = float("inf")
    else:
        snap_stop = snap_until_s if snap_end is None else snap_end
    sn = np.arange(N + 1) if snap_nodes is None else np.asarray(snap_nodes, int)
    snap_t, snap_T, snap_C, snap_R = [], [], [], []
    if snap_every_s > 0:
        snap_next = (math.floor(t / snap_every_s + 1e-12) + 1) * snap_every_s
        if t == 0.0 and snap_stop >= 0.0:
            snap_t.append(0.0); snap_T.append(T[sn].copy())
            snap_C.append(C[sn].copy()); snap_R.append(R_now)
    else:
        snap_next = None

    tr = {k: (Grow() if cmax_track else None)
          for k in ("t", "cmax", "arg", "c0", "cR", "cmin")}
    ckpt, profiles = {}, {}
    prof_left = sorted(float(v) for v in profile_times)
    picard_iters = Grow()
    n_not_conv = 0
    res_max_T = res_max_C = 0.0
    dT_seen = dC_seen = 0.0

    aud = dict(n=0, max_abs=0.0, max_rel_global=None, max_rel_perstep=0.0,
               max_rel_env=0.0, sum_abs=0.0,
               samples=[], peak_dE=0.0, peak_flux=0.0, n_ge_1e8=0,
               max_rel_transient=0.0, n_transient=0, max_abs_transient=0.0,
               max_rel_env_transient=0.0, E_scale_max=0.0)
    om_n = omega_of_C(C, P)
    energy_terms = []
    mass_aud = dict(max_rel=0.0, n_ge_1e8=0, max_abs=0.0, worst_t=None, first=None)
    aud_s = (dict(t=Grow(), lhs=Grow(), rhs=Grow(), rel=Grow(), cmax=Grow())
             if audit_series else None)
    m_prev = float(np.dot(W, C))
    flux_cum_audit = 0.0
    flux_cum_phys = 0.0

    dt_now = dt1
    n_switch = 0
    t_end = None
    t_dry = None
    t_dry_prev = None
    cmax_prev = None
    status = "unknown"
    n_steps_total = step
    mass = m_prev
    R_min_seen, R_max_seen = R_now, R_now
    g_geo_prev = None

    while True:
        if stop_on_dry:
            if t_dry is not None and t >= t_end - 1e-9:
                status = "ok"
                break
        else:
            if t_end_override is not None and t >= t_end_override - 1e-9:
                status = "fixed_end"
                break
        if t >= t_cap - 1e-9:
            status = "not_dry_by_cap"
            break
        if n_switch == 0 and dt2 != dt1 and t >= t_switch - 1e-12:
            dt_now = dt2
            n_switch = 1
            if verbose:
                print("[stage] 段切换：t = %.6f s  Δt = %g s" % (t, dt_now), flush=True)

        t_new = t + dt_now
        ta_new, ca_new = drv_at(t_new)
        R_new = float(radius_at(t_new, R_mode))
        R_old = R_now
        if R_mode in ("shrink", "pchip"):
            if R_new > R_old + 1e-15:
                raise AssertionError("R(t) 非单调不增（t=%.6f：%r  %r）"
                                     % (t_new, R_old, R_new))
            if not (R_MIN - 1e-12 <= R_new <= R0 + 1e-12):
                raise AssertionError("R(t) 越界（t=%.6f：R=%r）" % (t_new, R_new))
        R_min_seen = min(R_min_seen, R_new)
        R_max_seen = max(R_max_seen, R_new)

        R_geo = R_new if geo_layer == "nplus1" else R_old
        g = dt_now / (R_geo * R_geo)

        k_it = 0
        converged = False
        C_it = C
        T_it = np.full(N + 1, ta_new) if T_mode == "prescribed" else T
        C_prev, T_prev = C_it, T_it
        mT = None
        mC = None
        for m in range(1, picard_max + 1):
            k_it = m
            om = omega_of_C(C_it, P)
            kv = k_of_C(C_it, P)
            Dv = eval_D(logD_of_CT(C_it, T_it, P), floor_mon, C_it, T_it, t_new)

            if T_mode == "coupled" and block_order == "TC":
                kf = face_gam("k", kv, C_it, T_it, face_rule, P)
                T_it, mT = assemble(g, geo, kf, om, P["H"], R_new, bc, T, ta_new)
                Dv = eval_D(logD_of_CT(C_it, T_it, P), floor_mon, C_it, T_it, t_new)

            Df = face_gam("D", Dv, C_it, T_it, face_rule, P)
            C_it, mC = assemble(g, geo, Df, np.ones(N + 1), P["HM"], R_new, bc, C, ca_new)

            if T_mode == "coupled" and block_order == "CT":
                om = omega_of_C(C_it, P)
                kv = k_of_C(C_it, P)
                kf = face_gam("k", kv, C_it, T_it, face_rule, P)
                T_it, mT = assemble(g, geo, kf, om, P["H"], R_new, bc, T, ta_new)

            dT = float(np.abs(T_it - T_prev).max())
            dC = float(np.abs(C_it - C_prev).max())
            T_prev, C_prev = T_it, C_it
            dT_seen = max(dT_seen, dT)
            dC_seen = max(dC_seen, dC)
            if T_mode == "prescribed":
                if dC < picard_tol:
                    converged = True
                    break
            elif dT < picard_tol and dC < picard_tol:
                converged = True
                break

        if T_mode == "coupled":
            res_max_T = max(res_max_T, resid_of(mT, T_it))
        res_max_C = max(res_max_C, resid_of(mC, C_it))

        picard_iters.append(k_it)
        if not converged:
            n_not_conv += 1

        T_new, C_new = T_it, C_it
        cmax_now = float(C_new.max())
        cmin_now = float(C_new.min())
        tmin_now = float(T_new.min())
        tmax_now = float(T_new.max())

        if cmin_now < c_lo - 1e-12 or cmax_now > c_hi + 1e-12:
            viol.append(dict(t_s=float(t_new), kind="C", min=cmin_now, max=cmax_now,
                             bound=[c_lo, c_hi]))
        if tmin_now < t_lo - 1e-9 or tmax_now > t_hi + 1e-9:
            viol.append(dict(t_s=float(t_new), kind="T", min=tmin_now, max=tmax_now,
                             bound=[t_lo, t_hi]))
        if viol:
            status = "extremum_violation"
            t, n_steps_total = t_new, step + 1
            break

        if audit and bc == "robin":
            lhs_m = float(np.dot(W, C_new - C))
            rhs_m = dt_now * P["HM"] * (ca_new - float(C_new[N])) / R_new
            sc_m = max(abs(lhs_m), abs(rhs_m))
            rel_m = abs(lhs_m - rhs_m) / sc_m if sc_m > 0 else 0.0
            if rel_m > mass_aud["max_rel"]:
                mass_aud["max_rel"] = rel_m
                mass_aud["max_abs"] = abs(lhs_m - rhs_m)
                mass_aud["worst_t"] = float(t_new)
            if mass_aud["first"] is None:
                mass_aud["first"] = dict(t_s=float(t_new), lhs=lhs_m, rhs=rhs_m, rel=rel_m)
            if rel_m > 1e-8:
                mass_aud["n_ge_1e8"] += 1
            if aud_s is not None:
                aud_s["t"].append(float(t_new)); aud_s["lhs"].append(lhs_m)
                aud_s["rhs"].append(rhs_m); aud_s["rel"].append(rel_m)
                aud_s["cmax"].append(cmax_now)
            if rel_m > 1e-10:
                raise AssertionError(
                    "水分守恒审计越界（§6.6）：t=%.6f s 相对残差 %.3e > 1e-10 —— "
                    "先查累积项右端取的是哪一层（R3 条款），再查边界符号、再查系数层。"
                    % (t_new, rel_m))
            flux_cum_audit += dt_now * P["HM"] * (ca_new - float(C_new[N])) / R_new
            flux_cum_phys += dt_now * R_new * P["HM"] * (float(C_new[N]) - ca_new)

        om_new = omega_of_C(C_new, P)
        Tn = T if T_mode == "coupled" else np.full(N + 1, ta_old)
        Tp = T_new
        if audit:
            lhsA = float(np.dot(W, om_new * Tp) - np.dot(W, om_n * Tn)
                         - np.dot(W, Tn * (om_new - om_n)))
            rhsA = dt_now * P["H"] * (ta_new - float(Tp[N])) / R_new
            residA = lhsA - rhsA
            scA = max(abs(lhsA), abs(rhsA))
            relA = abs(residA) / scA if scA > 0 else 0.0
            E_scale = float(np.dot(W, np.abs(om_new * Tp)))
            aud["E_scale_max"] = max(aud["E_scale_max"], E_scale)
            relEnv = abs(residA) / E_scale if E_scale > 0 else 0.0
            aud["n"] += 1
            aud["max_abs"] = max(aud["max_abs"], abs(residA))
            aud["max_rel_perstep"] = max(aud["max_rel_perstep"], relA)
            aud["max_rel_env"] = max(aud["max_rel_env"], relEnv)
            aud["peak_dE"] = max(aud["peak_dE"], abs(lhsA))
            aud["peak_flux"] = max(aud["peak_flux"], abs(rhsA))
            aud["sum_abs"] += abs(residA)
            ref = relEnv
            if ref > 1e-8:
                aud["n_ge_1e8"] += 1
            if t_new <= T_FREEZE_S:
                aud["n_transient"] += 1
                aud["max_rel_transient"] = max(aud["max_rel_transient"], ref)
                aud["max_rel_env_transient"] = max(aud["max_rel_env_transient"], relEnv)
                aud["max_abs_transient"] = max(aud["max_abs_transient"], abs(residA))
            if len(aud["samples"]) < 5:
                aud["samples"].append(dict(t_s=float(t_new), lhs_A=lhsA, rhs_A=rhsA,
                                           residual=residA, rel_self=relA,
                                           rel_envelope=relEnv))
            if T_mode == "coupled" and relEnv > 1e-8:
                raise AssertionError(
                    "能量审计越界（§6.7 写法 A）：t=%.6f s 包络相对残差 %.3e > 1e-8 "
                    "（绝对 %.3e；自尺度 %.3e）—— 先查累积项冻结层（R3）、再查边界项符号、"
                    "再查 ω 取的是哪一层。" % (t_new, relEnv, abs(residA), relA))
        if audit and len(energy_terms) < 200000:
            dH_phys = (R_new ** 2) * float(np.dot(W, om_new * Tp)) - \
                      (R_old ** 2) * float(np.dot(W, om_n * Tn))
            cap_acct = (R_new ** 2) * float(np.dot(W, Tn * (om_new - om_n)))
            geo_term = (R_new ** 2 - R_old ** 2) * float(np.dot(W, om_n * Tn))
            bnd_flux = dt_now * R_new * P["H"] * (ta_new - float(Tp[N]))
            energy_terms.append((float(t_new), dH_phys, cap_acct, geo_term, bnd_flux))
        om_n = om_new
        ta_old = ta_new

        T, C = T_new, C_new
        R_now = R_new
        mass = float(np.dot(W, C_new))
        t, step = t_new, step + 1
        n_steps_total = step

        if cmax_track:
            tr["t"].append(t); tr["cmax"].append(cmax_now)
            tr["arg"].append(float(np.argmax(C_new)))
            tr["c0"].append(float(C_new[0])); tr["cR"].append(float(C_new[-1]))
            tr["cmin"].append(cmin_now)

        if stop_on_dry and t_dry is None and cmax_now <= c_dry:
            t_dry = t
            t_dry_prev = (t - dt_now, cmax_prev)
            t_end = t + margin_s
            n_end = math.ceil(t_end / dt_now - 1e-9)
            t_end = n_end * dt_now
            if verbose:
                print("[dry] 首达 C_max ≤ %g：t = %.6f s (%.4f h, %.4f d)；C_max=%.9f；"
                      " t_end = %.4f s" % (c_dry, t, t / 3600.0, t / 86400.0,
                                            cmax_now, t_end), flush=True)

        if snap_next is not None:
            while t >= snap_next - 1e-9 and snap_next <= snap_stop + 1e-9:
                snap_t.append(t); snap_T.append(T[sn].copy())
                snap_C.append(C[sn].copy()); snap_R.append(R_now)
                snap_next += snap_every_s

        while prof_left and t >= prof_left[0] - 1e-9:
            profiles[prof_left.pop(0)] = (T.copy(), C.copy())

        for tc in checkpoints:
            if tc not in ckpt and t >= tc - 1e-9:
                ckpt[tc] = dict(t=t, T=T.copy(), C=C.copy(), cum_flux=cum_flux,
                                mass=mass, n_steps=step, floor_count=floor_mon.count,
                                N=N, R_mode=R_mode)

        if verbose and progress_every and step % progress_every == 0:
            nn = min(2000, picard_iters.n)
            print("[prog] t=%10.3f s (%.3f h) n=%8d C_max=%.9f C(0)=%.9f R=%.6f cm "
                  "it=%.2f floor=%d wall=%.1f s"
                  % (t, t / 3600.0, step, cmax_now, float(C_new[0]), R_now * 100.0,
                     float(picard_iters.a[picard_iters.n - nn:picard_iters.n].mean()),
                     floor_mon.count, time.time() - t_wall0), flush=True)
        cmax_prev = cmax_now

    iters = picard_iters.array()
    uniq, cnt = (np.unique(iters, return_counts=True) if iters.size else
                 (np.array([]), np.array([])))
    et = np.array(energy_terms) if energy_terms else np.zeros((0, 5))
    return dict(
        status=status, params=P,
        xi=xi, dxi=dxi, N=N, W=W, R_mode=R_mode, geo_layer=geo_layer,
        face_rule=face_rule, props=props, bc=bc, T_mode=T_mode, block_order=block_order,
        d_const=d_const,
        dt1=dt1, dt2=dt2, t_switch=t_switch, dt_final=dt_now, n_switch=n_switch,
        n_steps=n_steps_total, t_final=t, t_dry=t_dry, t_dry_prev=t_dry_prev,
        t_end=t_end, c_dry=c_dry, R_min_seen=R_min_seen, R_max_seen=R_max_seen,
        snap_t=np.array(snap_t), snap_T=np.array(snap_T), snap_C=np.array(snap_C),
        snap_R=np.array(snap_R), snap_nodes=sn,
        T_final=T.copy(), C_final=C.copy(), R_final=R_now,
        picard=dict(max_iters=int(iters.max()) if iters.size else 0,
                    mean_iters=float(iters.mean()) if iters.size else 0.0,
                    median_iters=float(np.median(iters)) if iters.size else 0.0,
                    n_not_converged=int(n_not_conv),
                    hist_iters=[int(v) for v in uniq],
                    hist_counts=[int(v) for v in cnt]),
        residual=dict(max_T=res_max_T, max_C=res_max_C),
        picard_step_dmax=dict(T=dT_seen, C=dC_seen),
        floor=floor_mon.to_dict(),
        violations=viol,
        extremum_bounds=dict(T_C=[t_lo, t_hi], C_kgkg=[c_lo, c_hi]),
        audit=dict(n=aud["n"], max_abs=aud["max_abs"],
                   max_rel_global=(aud["max_abs"]
                                   / max(aud["peak_dE"], aud["peak_flux"], 1e-300)),
                   max_rel_perstep=aud["max_rel_perstep"],
                   max_rel_env=aud["max_rel_env"],
                   E_scale_max=aud["E_scale_max"],
                   n_ge_1e8=aud["n_ge_1e8"],
                   max_rel_transient=aud["max_rel_transient"],
                   max_rel_env_transient=aud["max_rel_env_transient"],
                   max_abs_transient=aud["max_abs_transient"],
                   n_transient=aud["n_transient"],
                   peak_dE=aud["peak_dE"], peak_flux=aud["peak_flux"],
                   mean_abs=(aud["sum_abs"] / aud["n"]) if aud["n"] else 0.0,
                   samples=aud["samples"],
                   form="写法 A（Q4-E-A，ξ 测度、无几何项、与 Q2-25 形式同构）",
                   applicable=(T_mode == "coupled" and bc == "robin"),
                   criterion_note=("判据 = `max_rel_global`（= 全程 max|残差| / "
                                   "全程 max(|ΔH|,|边界通量|)，与 Q2 的 §8.7 同构，"
                                   "≤1e-8）；逐走判据用 `max_rel_env`（除以焓包络 "
                                   "ΣWω|T|，≤1e-8），因小步上自尺度会退化。"),
                   note=("T_mode='coupled' 时 (Q4-E-A) 对本离散格式是代数恒等式"
                         "（判据 `max_rel_global` ≤1e-8，与 Q2 §8.7 同构）；"
                         "T_mode='prescribed'（主口径）下不求解能量方程，该式在恒温段"
                         "（t>14400 s，T_a 常数）两端同为零、残差应为舍入零，在升温段的"
                         "残差是「预设 T 口径的热失衡诊断量」。预设口径下自尺度退化"
                         "（两端可同时为零） 用 `max_rel_env`（除以焓包络 "
                         "`ΣWω|T|`）作为普适量度，不作为审计判据。")),
        audit_energy_terms=dict(
            note="写法 B（物理焓）三分项逐内部步（§6.7 必落盘）；几何项绝不进右端通量",
            header=["t_s", "dH_phys", "cap_account_R2", "geo_term", "boundary_flux"],
            n=int(et.shape[0]), values=et),
        mass=dict(m0=m_prev, m1=mass,
                  ratio=mass / m_prev if m_prev else float("nan"),
                  cum_flux=flux_cum_audit, cum_flux_phys=flux_cum_phys,
                  rel_err=(abs((mass - m_prev) - flux_cum_audit) / abs(mass - m_prev)
                           if m_prev != mass else 0.0),
                  note="P3（§6.6）：Δ(ΣW_i^ξ C_i) 应等于 Σ Δt·h_m(C_a−C_N)/R；"
                       "跨界面通量逐对相消，严格成立"),
        mass_audit=mass_aud,
        mass_audit_series=(None if aud_s is None else
                           dict(t=aud_s["t"].array(), lhs=aud_s["lhs"].array(),
                                rhs=aud_s["rhs"].array(), rel=aud_s["rel"].array(),
                                cmax=aud_s["cmax"].array())),
        cmax=dict(t=tr["t"].array() if cmax_track else None,
                  cmax=tr["cmax"].array() if cmax_track else None,
                  argmax=tr["arg"].array() if cmax_track else None,
                  c0=tr["c0"].array() if cmax_track else None,
                  cR=tr["cR"].array() if cmax_track else None,
                  cmin=tr["cmin"].array() if cmax_track else None),
        checkpoints=ckpt, profiles=profiles,
        elapsed_s=time.time() - t_wall0,
    )

def setup_params(N, R_mode, T_mode, bc, ca_scenario, const_ca, const_ta, d_const, dt1, dt2, geo_layer, overrides, props, snap_every_s, ta_scenario):
    t_wall0 = time.time()
    P = default_params()
    P.update(overrides)
    P["props"] = props
    P["D_CONST"] = d_const
    P["LOG_D0"] = math.log(P["D0"])
    P["LOG_A2_D0"] = math.log(P["A2_D0"])
    if R_mode not in ("shrink", "const", "pchip"):
        raise ValueError("R_mode 必须为 'shrink' / 'const' / 'pchip'")
    if geo_layer not in ("nplus1", "n"):
        raise ValueError("geo_layer 必须为 'nplus1' 或 'n'")
    if T_mode not in ("prescribed", "coupled"):
        raise ValueError("T_mode 必须为 'prescribed' 或 'coupled'")
    if bc not in ("robin", "dirichlet"):
        raise ValueError("bc 必须为 'robin' 或 'dirichlet'")
    for nm, dtc in (("dt1", dt1), ("dt2", dt2)):
        if dtc <= 0 or abs(round(60.0 / dtc) * dtc - 60.0) > 1e-9:
            raise ValueError("%s=%g 不能整除 60 s（§7.1 C4：输出必须落在内部步上）"
                             % (nm, dtc))
    for nm, sv in (("snap_every_s", snap_every_s),):
        if sv > 0:
            for dtc in (dt1, dt2):
                if abs(round(sv / dtc) * dtc - sv) > 1e-9:
                    raise ValueError("%s=%g 不是 Δt=%g 的整数倍" % (nm, sv, dtc))

    xi, dxi, N, W = build_grid(N)
    geo = make_geo(N, dxi, W)

    drv = driver_table()
    Ta_scen = scenario_driver(ta_scenario, drv["T"], const_ta)
    Ca_scen = scenario_driver(ca_scenario, drv["C"], const_ca)
    use_frozen_tbl = (ta_scenario == "freeze_last" and ca_scenario == "freeze_last")
    _inv_dt, _nseg = drv["inv_dt"], drv["nseg"]
    return Ca_scen, N, P, Ta_scen, W, _inv_dt, _nseg, dxi, geo, t_wall0, use_frozen_tbl, xi

def prepare_drivers(Ca_scen, P, Ta_scen, _inv_dt, _nseg, driver_interp, use_frozen_tbl):
    def drv_at(t):
        if use_frozen_tbl:
            return driver_at(t, driver_interp)
        if t >= T_FREEZE_S:
            return float(Ta_scen[-1]), float(Ca_scen[-1])
        x = t * _inv_dt
        i = int(x)
        if i >= _nseg:
            i = _nseg - 1
        f = x - i
        return (float(Ta_scen[i] + (Ta_scen[i + 1] - Ta_scen[i]) * f),
                float(Ca_scen[i] + (Ca_scen[i + 1] - Ca_scen[i]) * f))

    ta_bound = max(float(np.max(Ta_scen)), float(Ta_scen[-1]))
    ca_bound = min(float(np.min(Ca_scen)), float(Ca_scen[-1]))
    c_lo = min(float(P["C_INIT"]), ca_bound)
    c_hi = float(P["C_INIT"])
    t_lo = min(float(P["T_INIT"]), float(np.min(Ta_scen)))
    t_hi = max(float(P["T_INIT"]), ta_bound)
    return c_hi, c_lo, drv_at, t_hi, t_lo


def smoke_test(N=200, dt=20.0, t_end=21600.0, verbose=True):
    o = solve_q4(N=N, dt1=dt, dt2=dt, t_switch=1e18, stop_on_dry=False,
                 t_end_override=t_end, snap_every_s=60.0, snap_until_s=t_end,
                 cmax_track=True, verbose=False)
    C = o["C_final"]
    checks = {
        "C_max_decreasing": bool(C.max() < C_INIT),
        "C_radial_nonincreasing": bool(np.all(np.diff(C) <= 1e-12)),
        "C_below_init": bool(C.max() <= C_INIT + 1e-12),
        "C_above_ca": bool(C.min() >= CA_MIN - 1e-12),
        "no_floor_activation": bool(o["floor"]["count"] == 0),
        "picard_all_converged": bool(o["picard"]["n_not_converged"] == 0),
        "no_extremum_violation": bool(len(o["violations"]) == 0),
        "mass_audit_le_1e-8": bool(o["mass_audit"]["max_rel"] <= 1e-8),
        "R_within_bounds": bool(o["R_min_seen"] >= R_MIN - 1e-12),
    }
    if verbose:
        print("耗时 %.2f s；步数 %d；Picard 均值 %.3f 最大 %d 未收敛 %d"
              % (o["elapsed_s"], o["n_steps"], o["picard"]["mean_iters"],
                 o["picard"]["max_iters"], o["picard"]["n_not_converged"]))
        print("C: 中心 %.9f  表面 %.9f  C_max %.9f  C_min %.9f"
              % (C[0], C[-1], C.max(), C.min()))
        print("水分审计 max 相对残差 %.3e（worst t=%s）；D_floor 激活 %d 次"
              % (o["mass_audit"]["max_rel"], o["mass_audit"]["worst_t"],
                 o["floor"]["count"]))
        for k, v in checks.items():
            print("  %-24s %s" % (k, "PASS" if v else "FAIL"))
    return o, checks


if __name__ == "__main__":
    print("[solver_q4] smoke test  N=200, dt=20 s, t_end=21600 s")
    o, ck = smoke_test()
    print("全部自检通过" if all(ck.values()) else "存在失败项")
